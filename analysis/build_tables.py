from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "tables"


DEFINITIONS = [
    ("Nominal package quantity", "P", "Labelled or target package content"),
    ("Purchase at event t", "Qₜ", "Packaged quantity entering usable inventory"),
    ("Demand at event t", "Dₜ", "Quantity requested at a stochastic usage event"),
    ("Consumption at event t", "Cₜ", "Quantity served from usable inventory"),
    ("Inventory after event t", "Iₜ", "Quantity retained and carried forward"),
    ("Backlog after event t", "Bₜ", "Demand not yet served"),
    ("Inventory exit at event t", "Eₜ", "Spoilage, expiry, forced removal, or disposal"),
    ("Terminal residual", "R", "Diagnostic remainder retained as inventory"),
    (
        "Generalized cost",
        "J",
        "Monetary product, purchase, packaging, storage, freshness, handling, "
        "adaptation, and variety costs",
    ),
]


def definitions_table() -> pd.DataFrame:
    return pd.DataFrame(DEFINITIONS, columns=["Concept", "Symbol", "Operational definition"])


def calibration_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "tables/parameter_evidence_table.csv")
    return frame[
        [
            "parameter",
            "configured_value",
            "unit",
            "evidence_class",
            "model_role",
        ]
    ]


def weighted_replication_standard_errors(
    frame: pd.DataFrame,
    metric: str,
) -> pd.Series:
    weighted = frame.assign(
        weighted_metric=frame[metric] * frame["scenario_weight"]
    )
    totals = weighted.groupby(
        ["package_kg", "replication"], as_index=False
    )["weighted_metric"].sum()
    grouped = totals.groupby("package_kg")["weighted_metric"]
    return grouped.std(ddof=1) / np.sqrt(grouped.count())


def rice_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/baseline_size_comparison.csv")
    replications = pd.read_csv(ROOT / "output/simulation_replications.csv.gz")
    replications = replications[replications["regime"] == "baseline"]
    confidence_z = 1.959963984540054
    metrics = [
        "generalized_cost_jpy",
        "physical_discard_kg",
        "purchases",
        "packaging_mass_kg",
        "mean_age_at_consumption_days",
        "cost_regret_jpy",
    ]
    for metric in metrics:
        frame[f"weighted_{metric}"] = frame[metric] * frame["scenario_weight"]
    summary = (
        frame.groupby("package_kg")[[f"weighted_{metric}" for metric in metrics]]
        .sum()
        .rename(columns=lambda value: value.replace("weighted_", ""))
        .reset_index()
    )
    for metric in ["generalized_cost_jpy", "physical_discard_kg"]:
        standard_error = (
            summary["package_kg"]
            .map(weighted_replication_standard_errors(replications, metric))
            .to_numpy()
        )
        summary[f"{metric}_ci_low"] = np.maximum(
            0.0, summary[metric] - confidence_z * standard_error
        )
        summary[f"{metric}_ci_high"] = (
            summary[metric] + confidence_z * standard_error
        )
    summary["generalized_cost_jpy_95_ci"] = summary.apply(
        lambda row: (
            f"{row['generalized_cost_jpy']:.1f} "
            f"({row['generalized_cost_jpy_ci_low']:.1f}–"
            f"{row['generalized_cost_jpy_ci_high']:.1f})"
        ),
        axis=1,
    )
    summary["physical_discard_kg_95_ci"] = summary.apply(
        lambda row: (
            f"{row['physical_discard_kg']:.3f} "
            f"({row['physical_discard_kg_ci_low']:.3f}–"
            f"{row['physical_discard_kg_ci_high']:.3f})"
        ),
        axis=1,
    )
    return summary[
        [
            "package_kg",
            "generalized_cost_jpy_95_ci",
            "physical_discard_kg_95_ci",
            "purchases",
            "packaging_mass_kg",
            "mean_age_at_consumption_days",
            "cost_regret_jpy",
        ]
    ].round(3)


def cross_product_table() -> pd.DataFrame:
    return pd.read_csv(ROOT / "data/processed/cross_product_classification.csv")


def negative_control_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/mechanism_pathway_comparison.csv")
    selected = frame[
        frame["comparison_basis"] == "baseline_best_single"
    ].copy()
    labels = {
        "baseline": "Discrete + spoilage",
        "no_discrete_consumption_quantum": "Lognormal + spoilage",
        "randomized_non_lattice_usage": "Gamma + spoilage",
        "full_carry_no_spoilage": "Discrete + no spoilage",
        "continuous_no_spoilage": "Lognormal + no spoilage",
    }
    selected["control"] = selected["regime"].map(labels)
    return selected[
        [
            "control",
            "package_kg",
            "generalized_cost_jpy",
            "physical_discard_kg",
            "purchases",
            "mean_age_at_consumption_days",
            "packaging_mass_kg",
        ]
    ].round(3)


def sku_threshold_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/sku_inclusion_thresholds.csv")
    baseline = frame[frame["regime"] == "baseline"].copy()
    return baseline[
        [
            "from_sku_count",
            "to_sku_count",
            "best_from_portfolio",
            "best_to_portfolio",
            "from_consumer_cost_jpy",
            "to_consumer_cost_jpy",
            "break_even_complexity_jpy_per_added_sku",
            "globally_optimal_portfolios_nested",
        ]
    ].round(3)


def join_values(values: list) -> str:
    return "/".join(f"{value:g}" for value in values)


def household_table() -> pd.DataFrame:
    config = yaml.safe_load((ROOT / "config/default.yaml").read_text())
    rows = [
        {
            "scenario": household["name"],
            "weight": round(household["scenario_weight"], 5),
            "annual_consumption_kg": household["annual_consumption_kg"],
            "usage_quanta_g": join_values(household["usage_g"]),
            "usage_probabilities": join_values(household["usage_probability"]),
            "lattice_spacing_g": math.gcd(*household["usage_g"]),
            "adaptation_probability": household["adaptation_probability"],
            "adaptation_min_ratio": household["adaptation_min_ratio"],
            "freshness_threshold_days": household["freshness_threshold_days"],
            "handling_threshold_kg": household["handling_threshold_kg"],
        }
        for household in config["households"]
    ]
    return pd.DataFrame(rows)


def regime_table() -> pd.DataFrame:
    config = yaml.safe_load((ROOT / "config/default.yaml").read_text())
    rows = []
    for regime in config["regimes"]:
        changes = [
            f"{key.replace('_', ' ')} = {value:g}"
            if isinstance(value, (int, float))
            else f"{key.replace('_', ' ')} = {value}"
            for key, value in regime.items()
            if key != "name"
        ]
        rows.append(
            {
                "regime": regime["name"],
                "modification_from_baseline": "; ".join(changes)
                or "anchor calibration",
            }
        )
    return pd.DataFrame(rows)


def factorial_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/lattice_factorial_decomposition.csv")
    components = [
        "lattice_main_at_no_spoilage",
        "spoilage_main_under_continuous",
        "lattice_spoilage_interaction",
    ]
    cost = frame[frame["metric"] == "generalized_cost_jpy"].set_index("package_kg")
    discard = frame[frame["metric"] == "physical_discard_kg"].set_index(
        "package_kg"
    )
    table = pd.DataFrame({"package_kg": cost.index})
    for component in components:
        label = component.split("_")[0]
        if component.endswith("interaction"):
            label = "interaction"
        table[f"cost_{label}_jpy"] = cost[component].round(1).to_numpy()
    for component in components:
        label = component.split("_")[0]
        if component.endswith("interaction"):
            label = "interaction"
        table[f"discard_{label}_kg"] = (
            discard.loc[cost.index, component].round(4).to_numpy() + 0.0
        )
    return table


def consumption_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/continuous_consumption_summary.csv")
    return pd.DataFrame(
        {
            "consumption_multiplier": frame["annual_consumption_multiplier"],
            "scenario_range_kg": [
                f"{low:.1f}–{high:.1f}"
                for low, high in zip(
                    frame["minimum_scenario_annual_consumption_kg"],
                    frame["maximum_scenario_annual_consumption_kg"],
                )
            ],
            "best_one_sku": frame["best_one_sku_portfolio"],
            "best_two_sku": frame["best_two_sku_portfolio"],
            "best_three_sku": frame["best_three_sku_portfolio"],
            "global_portfolio": frame["global_optimal_portfolio"],
            "one_to_two_threshold_jpy": frame[
                "one_to_two_break_even_complexity_jpy"
            ].round(1),
            "two_to_three_threshold_jpy": frame[
                "two_to_three_break_even_complexity_jpy"
            ].round(1),
            "benchmark_regret_jpy": frame["benchmark_regret_jpy"].round(1),
        }
    )


def phase_threshold_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/sku_phase_threshold_curves.csv")
    keys = ["spoilage_hazard_multiplier", "purchase_event_cost_multiplier"]
    first = frame[frame["from_sku_count"] == 1].set_index(keys)
    second = frame[frame["from_sku_count"] == 2].set_index(keys)
    table = pd.DataFrame(
        {
            "best_one_sku": first["best_from_portfolio"],
            "best_two_sku": first["best_to_portfolio"],
            "one_to_two_threshold_jpy": first[
                "break_even_complexity_jpy_per_added_sku"
            ].round(1),
            "best_three_sku": second.loc[first.index, "best_to_portfolio"],
            "two_to_three_threshold_jpy": second.loc[
                first.index, "break_even_complexity_jpy_per_added_sku"
            ].round(1),
            "one_two_nested": first["globally_optimal_portfolios_nested"],
        }
    )
    return table.reset_index().rename(
        columns={
            "spoilage_hazard_multiplier": "spoilage_multiplier",
            "purchase_event_cost_multiplier": "purchase_cost_multiplier",
        }
    )


def candidate_grid_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/candidate_grid_robustness.csv")
    return frame.round(
        {
            "system_cost_jpy": 1,
            "consumer_cost_jpy": 1,
            "physical_discard_kg": 3,
            "packaging_mass_kg": 3,
        }
    )


def regret_table() -> pd.DataFrame:
    config = yaml.safe_load((ROOT / "config/default.yaml").read_text())
    regimes = [regime["name"] for regime in config["regimes"]]
    frame = pd.read_csv(ROOT / "output/robust_portfolios.csv")
    frame = frame.sort_values(["maximum_regret_jpy", "mean_regret_jpy"]).head(10)
    return pd.DataFrame(
        {
            "sku_count": frame["sku_count"],
            "portfolio": frame["portfolio"],
            "maximum_regret_jpy": frame["maximum_regret_jpy"].round(1),
            "mean_regret_jpy": frame["mean_regret_jpy"].round(1),
            "baseline_regret_jpy": frame["baseline"].round(1),
            "regime_of_maximum": frame[regimes].idxmax(axis=1),
        }
    )


def manufacturing_table() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/manufacturing_precision.csv")
    return pd.DataFrame(
        {
            "target_kg": frame["target_g"] / 1000.0,
            "process_sigma_g": frame["process_sigma_g"],
            "maximum_underfill_probability": frame[
                "maximum_underfill_probability"
            ],
            "safety_margin_g": frame["safety_margin_g"].round(1),
            "mean_fill_g": frame["mean_fill_g"].round(1),
            "expected_giveaway_percent": (
                100 * frame["expected_giveaway_fraction"]
            ).round(3),
        }
    )


def add_dataframe(document: Document, title: str, frame: pd.DataFrame) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(title)
    run.bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, column in enumerate(frame.columns):
        table.rows[0].cells[index].text = str(column)
    for values in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    document.add_paragraph()


def main() -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    tables = [
        ("Table 1. Definitions and estimands", definitions_table()),
        ("Table 2. Parameter inputs and evidence status", calibration_table()),
        ("Table 3. Household design scenarios", household_table()),
        ("Table 4. Boundary and sensitivity regimes", regime_table()),
        ("Table 5. Matched lattice and exit-mechanism controls", negative_control_table()),
        ("Table 6. Lattice–spoilage factorial decomposition", factorial_table()),
        ("Table 7. Baseline single-size comparison", rice_table()),
        ("Table 8. Annual-consumption sensitivity", consumption_table()),
        ("Table 9. Global SKU-inclusion thresholds", sku_threshold_table()),
        ("Table 10. Dense SKU phase thresholds", phase_threshold_table()),
        ("Table 11. Candidate-grid robustness", candidate_grid_table()),
        ("Table 12. Leading scenario-set minimax-regret portfolios", regret_table()),
        ("Table 13. Fill-process precision scenarios", manufacturing_table()),
        ("Table 14. Cross-product mechanism hypotheses", cross_product_table()),
    ]
    for stale in TABLES.glob("table_*.csv"):
        stale.unlink()
    for index, (_, frame) in enumerate(tables, start=1):
        frame.to_csv(TABLES / f"table_{index}.csv", index=False)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8)
    document.add_heading("Editable tables: Package–Consumption Unit Mismatch", level=0)
    for title, frame in tables:
        add_dataframe(document, title, frame)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.save(TABLES / "pcum_tables_editable.docx")


if __name__ == "__main__":
    main()
