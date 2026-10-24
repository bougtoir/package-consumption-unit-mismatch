from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
import yaml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "submission" / "SMPT"
PUBLIC_REPOSITORY = "https://github.com/bougtoir/package-consumption-unit-mismatch"
TITLE = (
    "Coupling stochastic household-inventory simulation to package-SKU "
    "portfolio selection: an ablation-verified decision model"
)

FIGURES = {
    "concept": (
        1,
        "figure_1_conceptual_distinction.png",
        "Conceptual distinction among package purchasing, served consumption, carried "
        "inventory, carry-over or adaptation, and physical inventory exit. Residual "
        "inventory becomes discard only through an explicit exit mechanism.",
    ),
    "lattice": (
        2,
        "figure_2_lattice_residual.png",
        "Deterministic one-package lattice residual P mod 150 g across nominal package "
        "sizes. Orange markers are the ten coarse candidate sizes; the periodic "
        "remainder is an inventory state, not a discard measure.",
    ),
    "controls": (
        3,
        "figure_3_lattice_negative_control.png",
        "Matched discrete and continuous usage controls. Full carry-over without "
        "spoilage produces zero physical discard under both usage processes; "
        "discreteness alone is not an inventory-exit mechanism.",
    ),
    "consumption": (
        4,
        "figure_4_consumption_sensitivity.png",
        "Continuous annual-consumption sensitivity. The optimal single size changes "
        "with throughput, and the regret of the 2|5|10 kg benchmark is "
        "scenario-dependent.",
    ),
    "pareto": (
        5,
        "figure_5_pareto_frontier.png",
        "Generalized system cost and packaging mass for all one-, two-, and three-SKU "
        "portfolios in the baseline scenario. The black line identifies the "
        "non-dominated envelope in these two reported dimensions.",
    ),
    "break_even": (
        6,
        "figure_6_sku_break_even.png",
        "Global one-, two-, and three-SKU system costs as the linear complexity "
        "charge varies. Crossings are the analytical SKU-inclusion thresholds.",
    ),
    "phase": (
        7,
        "figure_7_sku_phase_map.png",
        "Optimal SKU-count phase map across purchase-event cost, spoilage intensity, "
        "and assortment-complexity cost. The colorbar categories are one, two, and "
        "three SKUs.",
    ),
    "grid": (
        8,
        "figure_8_candidate_grid_robustness.png",
        "Candidate-grid robustness of the scenario-weighted single-size objective. "
        "The blue curve is the exhaustively evaluated dense grid; orange markers show "
        "the original ten-size grid.",
    ),
    "manufacturing": (
        9,
        "figure_9_manufacturing_precision.png",
        "Fill-process safety margin above the nominal target and expected relative "
        "giveaway by target size under the analytical normal-process scenario. "
        "Smaller packages bear a larger relative giveaway.",
    ),
}

TABLES = {
    "definitions": (1, "Definitions and estimands.", ""),
    "parameters": (
        2,
        "Parameter inputs and evidence status.",
        "Note: empirical anchors and scenario assumptions are intentionally separated.",
    ),
    "households": (
        3,
        "Household design scenarios.",
        "Note: weights preserve only the census split between one-person and other "
        "households; lattice spacing is the greatest common divisor of the usage "
        "quanta.",
    ),
    "regimes": (
        4,
        "Boundary and sensitivity regimes.",
        "Note: parameters not listed retain their baseline values.",
    ),
    "controls": (
        5,
        "Matched lattice and inventory-exit controls at the baseline best single size.",
        "Note: confidence intervals represent Monte Carlo error conditional on the "
        "scenario parameters.",
    ),
    "factorial": (
        6,
        "Lattice–spoilage factorial decomposition of generalized cost and physical "
        "discard by package size.",
        "Note: components are differences in scenario-weighted means per "
        "household-year; they are accounting pathways, not independent causal effects.",
    ),
    "package_comparison": (
        7,
        "Scenario-weighted baseline comparison of single package sizes.",
        "Note: values are per household-year; intervals represent Monte Carlo error "
        "conditional on the scenario parameters; regret is relative to "
        "household-specific optimal sizes.",
    ),
    "consumption": (
        8,
        "Annual-consumption sensitivity of package and portfolio optima.",
        "Note: package identities are scenario outputs; the 2|5|10 kg lineup is a "
        "manufacturer example, not a market-share estimate.",
    ),
    "thresholds": (
        9,
        "Global SKU-inclusion thresholds under the baseline calibration.",
        "Note: thresholds compare globally optimal portfolios and do not assume "
        "nested assortments.",
    ),
    "phase_thresholds": (
        10,
        "Global SKU-inclusion thresholds across the dense spoilage–purchase grid.",
        "Note: thresholds are scenario break-even charges in JPY/household-year per "
        "added SKU, not observed producer costs.",
    ),
    "grid": (
        11,
        "Coarse- and dense-grid portfolio optima.",
        "Note: system cost adds the scenario complexity charge per added SKU to "
        "consumer cost.",
    ),
    "regret": (
        12,
        "Ten leading portfolios by maximum scenario-set regret.",
        "Note: regret is measured within each prespecified regime; the maximum does "
        "not represent distributional robustness.",
    ),
    "manufacturing": (
        13,
        "Fill-process precision scenarios.",
        "Note: the underfill probability is an analytical scenario constraint, not a "
        "legal standard.",
    ),
    "cross_product": (
        14,
        "Cross-product mechanism hypotheses and evidence limits.",
        "Note: entries are mechanism hypotheses, not empirical category findings.",
    ),
    "greedy": (
        15,
        "Nested greedy versus globally optimal portfolio selection by regime.",
        "Note: greedy retains the best k−1 set and adds the argmin next SKU; "
        "regret is the system-cost gap to the global optimum at that cardinality.",
    ),
    "deterministic": (
        16,
        "Deterministic mean-trajectory counterpart ablation of portfolio "
        "selection.",
        "Note: the deterministic counterpart uses constant mean usage and "
        "inter-purchase interval with no adaptive consumption; the regret is "
        "the cost of the deterministically chosen portfolio re-evaluated "
        "under the stochastic model.",
    ),
}

OBJECT_TOKEN = re.compile(r"<<(fig|tab):(\w+)>>")

# Items moved to the Supplement under the CAIE density triage (Gate 1):
# secondary scenario/regime detail, dense-grid/threshold enumerations,
# minimax portfolio ranking, manufacturing precision, and the cross-product
# illustration stay available as Tables S1-S9 and Figures S1-S4.
SUPPLEMENT_FIGURES = {"lattice", "pareto", "grid", "manufacturing"}
SUPPLEMENT_TABLES = {
    "households",
    "regimes",
    "factorial",
    "consumption",
    "phase_thresholds",
    "grid",
    "regret",
    "manufacturing",
    "cross_product",
}

REFERENCES = [
    (
        "Brinker, J., & Gündüz, H. I. (2016). Optimization of demand-related "
        "packaging sizes using a p-median approach. The International Journal "
        "of Advanced Manufacturing Technology, 87(5–8), 2259–2268. "
        "https://doi.org/10.1007/s00170-016-8630-4"
    ),
    (
        "Chan, R. B. Y. (2022). A review of packaging-related studies in the "
        "context of household food waste: Drivers, solutions and avenues for "
        "future research. Packaging Technology and Science, 35(1), 3–51. "
        "https://doi.org/10.1002/pts.2611"
    ),
    (
        "Fisher, M. L., & Ittner, C. D. (1999). The impact of product variety "
        "on automobile assembly operations: Empirical evidence and simulation "
        "analysis. Management Science, 45(6), 771–786. "
        "https://doi.org/10.1287/mnsc.45.6.771"
    ),
    (
        "Fresán, U., Errendal, S., Craig, W. J., & Sabaté, J. (2019). Does the "
        "size matter? A comparative analysis of the environmental impact of "
        "several packaged foods. Science of the Total Environment, 687, "
        "369–379. https://doi.org/10.1016/j.scitotenv.2019.06.109"
    ),
    (
        "Gaur, V., & Honhon, D. (2006). Assortment planning and inventory "
        "decisions under a locational choice model. Management Science, "
        "52(10), 1528–1543. https://doi.org/10.1287/mnsc.1060.0580"
    ),
    (
        "Gámez Albán, H. M., Soto Cardona, O. C., Mejía Argueta, C., & "
        "Sarmiento, A. T. (2015). A cost-efficient method to optimize package "
        "size in emerging markets. European Journal of Operational Research, "
        "241(3), 917–926. https://doi.org/10.1016/j.ejor.2014.09.020"
    ),
    (
        "Kandemir, C., Reynolds, C., Tom, Q., Fisher, K., Devine, R., "
        "Herszenhorn, E., Koh, S. C. L., & Evans, D. (2020). Using discrete "
        "event simulation to explore food wasted in the home. Journal of "
        "Simulation, 16(4), 415–435. "
        "https://doi.org/10.1080/17477778.2020.1829515"
    ),
    (
        "Kekre, S., & Srinivasan, K. (1990). Broader product line: A necessity "
        "to achieve success? Management Science, 36(10), 1216–1232. "
        "https://doi.org/10.1287/mnsc.36.10.1216"
    ),
    (
        "K\u00f6k, A. G., & Fisher, M. L. (2007). Demand estimation and "
        "assortment optimization under substitution: Methodology and "
        "application. Operations Research, 55(6), 1001–1021. "
        "https://doi.org/10.1287/opre.1070.0409"
    ),
    (
        "Boz, Z., Kiker, G., Haase, H., Orr, R., Vignesh, A., Campbell, C., "
        "Cohen, N., Gusto, C., Robertson, K., & Clemen, T. (2026). IFWASTE: "
        "An agent-based model for the estimation of household food waste. "
        "Sustainability, 18(12), 6091. https://doi.org/10.3390/su18126091"
    ),
    (
        "Koenigsberg, O., Kohli, R., & Montoya, R. (2010). Package size "
        "decisions. Management Science, 56(3), 485–494. "
        "https://doi.org/10.1287/mnsc.1090.1119"
    ),
    (
        "Malinowski, E., Karwan, M. H., Sun, L., & Pinto, J. M. (2018). "
        "Packaged gas supply chain planning with network-wide SKU "
        "rationalization. International Journal of Production Economics, 204, "
        "346–357. https://doi.org/10.1016/j.ijpe.2018.08.013"
    ),
    (
        "Menezes, M. B. C., Jalali, H., & Lamas, A. (2021). One too many: "
        "Product proliferation and the financial performance in manufacturing. "
        "International Journal of Production Economics, 242, 108285. "
        "https://doi.org/10.1016/j.ijpe.2021.108285"
    ),
    (
        "Moubed, M., Boroumandzad, Y., & Nadizadeh, A. (2021). A dynamic "
        "model for deteriorating products in a closed-loop supply chain. "
        "Simulation Modelling Practice and Theory, 108, 102269. "
        "https://doi.org/10.1016/j.simpat.2021.102269"
    ),
    (
        "Ministry of Agriculture, Forestry and Fisheries of Japan. (2026a). "
        "How much does rice weigh before and after cooking? "
        "https://www.maff.go.jp/j/heya/kodomo_sodan/0103/01.html "
        "(accessed 24 September 2026)"
    ),
    (
        "Ministry of Agriculture, Forestry and Fisheries of Japan. (2026b). "
        "How long can rice be stored? "
        "https://www.maff.go.jp/j/heya/kodomo_sodan/0006/01.html "
        "(accessed 24 September 2026)"
    ),
    (
        "Ministry of Agriculture, Forestry and Fisheries of Japan. (2026c). "
        "Rice consumption per person. "
        "https://www.maff.go.jp/j/heya/sodan/1808/01.html "
        "(accessed 24 September 2026)"
    ),
    (
        "Mino Rice Co., Ltd. (2026). Product lineup. "
        "https://minorice.jp/product/ (accessed 24 September 2026)"
    ),
    (
        "Nisshin Seifun Welna Inc. (2026). Ma-Ma spaghetti product information, "
        "GTIN 4902110304022. "
        "https://www.nisshin-seifun-welna.com/index/products/4902110304022.html "
        "(accessed 24 September 2026)"
    ),
    (
        "Omar, F., & de Silva, C. W. (2000a). Optimal portion control of "
        "natural objects with application in automated cannery processing of "
        "fish. Journal of Food Engineering, 46(1), 31–41. "
        "https://doi.org/10.1016/S0260-8774(00)00068-6"
    ),
    (
        "Omar, F., & de Silva, C. W. (2000b). The application of portion "
        "control optimization in an automated can-filling process. Journal of "
        "Food Process Engineering, 23(2), 101–125. "
        "https://doi.org/10.1111/j.1745-4530.2000.tb00506.x"
    ),
    (
        "Onggo, B. S. S., Panadero, J., Corlu, C. G., & Juan, A. A. (2019). "
        "Agri-food supply chains with stochastic demands: A multi-period "
        "inventory routing problem with perishable products. Simulation "
        "Modelling Practice and Theory, 97, 101970. "
        "https://doi.org/10.1016/j.simpat.2019.101970"
    ),
    (
        "Pauer, E., Wohner, B., Heinrich, V., & Tacker, M. (2019). Assessing "
        "the environmental sustainability of food packaging: An extended life "
        "cycle assessment including packaging-related food losses and waste "
        "and circularity assessment. Sustainability, 11(3), 925. "
        "https://doi.org/10.3390/su11030925"
    ),
    (
        "Randall, T., & Ulrich, K. (2001). Product variety, supply chain "
        "structure, and firm performance: Analysis of the U.S. bicycle "
        "industry. Management Science, 47(12), 1588–1604. "
        "https://doi.org/10.1287/mnsc.47.12.1588.10237"
    ),
    (
        "Statistics Bureau of Japan. (2025). Statistical Handbook of Japan "
        "2025. https://www.stat.go.jp/english/data/handbook/pdf/2025all.pdf"
    ),
    (
        "van Rooijen, M. A., Gerdessen, J. C., Claassen, G. D. H., & de Leeuw, "
        "S. L. J. M. (2024). Optimizing household food waste: The impact of "
        "meal planning, package sizes, and performance indicators. Resources, "
        "Conservation and Recycling, 205, 107559. "
        "https://doi.org/10.1016/j.resconrec.2024.107559"
    ),
    (
        "van Ryzin, G., & Mahajan, S. (1999). On the relationship between "
        "inventory costs and variety benefits in retail assortments. "
        "Management Science, 45(11), 1496–1509. "
        "https://doi.org/10.1287/mnsc.45.11.1496"
    ),
    (
        "Wan, X., Dresner, M. E., & Evers, P. T. (2014). Assessing the "
        "dimensions of product variety on performance: The value of product "
        "line and pack size. Journal of Business Logistics, 35(3), 213–224. "
        "https://doi.org/10.1111/jbl.12054"
    ),
    (
        "Wan, X., & Sanders, N. R. (2017). The negative impact of product "
        "variety: Forecast bias, inventory levels, and the role of vertical "
        "integration. International Journal of Production Economics, 186, "
        "123–131. https://doi.org/10.1016/j.ijpe.2017.02.002"
    ),
    (
        "Wensing, T., Sternbeck, M. G., & Kuhn, H. (2018). Optimizing "
        "case-pack sizes in the bricks-and-mortar retail trade. OR Spectrum, "
        "40(4), 913–944. https://doi.org/10.1007/s00291-018-0515-5"
    ),
    (
        "Williams, H., Lindström, A., Trischler, J., Wikström, F., & Rowe, Z. "
        "(2020). Avoiding food becoming waste in households—The role of "
        "packaging in consumers' practices across different food categories. "
        "Journal of Cleaner Production, 265, 121775. "
        "https://doi.org/10.1016/j.jclepro.2020.121775"
    ),
    (
        "Wilson, N. L. W., Rickard, B. J., Saputo, R., & Ho, S.-T. (2017). "
        "Food waste: The role of date labels, package size, and product "
        "category. Food Quality and Preference, 55, 35–44. "
        "https://doi.org/10.1016/j.foodqual.2016.08.004"
    ),
    (
        "Zhang, Y., & Qi, D. (2020). How households waste food at home: "
        "Estimating household food waste in a dynamic decision model under "
        "uncertainty. 2020 Annual Meeting, July 26–28, Kansas City, Missouri, "
        "Agricultural and Applied Economics Association, 304631. "
        "https://doi.org/10.22004/ag.econ.304631"
    ),
    (
        "Zhao, Q., Tan, K., Du, J., & Van Woensel, T. (2023). Joint case pack "
        "size and unpacking location optimization in a retail supply chain "
        "including product returns. Computers & Industrial Engineering, 182, "
        "109415. https://doi.org/10.1016/j.cie.2023.109415"
    ),
]


def load_inputs() -> dict:
    config = yaml.safe_load((ROOT / "config/default.yaml").read_text())
    frames = {
        f"table_{number}": pd.read_csv(ROOT / f"tables/table_{number}.csv")
        for number, _, _ in TABLES.values()
    }
    frames.update(
        {
            "size": pd.read_csv(ROOT / "output/size_household_metrics.csv"),
            "portfolio": pd.read_csv(ROOT / "output/portfolio_results.csv"),
            "contrasts": pd.read_csv(ROOT / "output/portfolio_contrasts.csv"),
            "robust": pd.read_csv(ROOT / "output/robust_portfolios.csv"),
            "terminal": pd.read_csv(ROOT / "output/terminal_residual.csv"),
            "manufacturing": pd.read_csv(
                ROOT / "output/manufacturing_precision.csv"
            ),
            "convergence": pd.read_csv(
                ROOT / "output/convergence_diagnostics.csv"
            ),
            "thresholds": pd.read_csv(
                ROOT / "output/sku_inclusion_thresholds.csv"
            ),
            "phase": pd.read_csv(ROOT / "output/sku_phase_map.csv"),
            "consumption": pd.read_csv(
                ROOT / "output/continuous_consumption_summary.csv"
            ),
            "consumption_transitions": pd.read_csv(
                ROOT / "output/consumption_transition_thresholds.csv"
            ),
            "grid": pd.read_csv(ROOT / "output/candidate_grid_robustness.csv"),
            "greedy": pd.read_csv(
                ROOT / "output/greedy_vs_global_portfolio.csv"
            ),
            "det_size": pd.read_csv(
                ROOT / "output/deterministic_size_household_metrics.csv"
            ),
            "det_compare": pd.read_csv(
                ROOT / "output/deterministic_counterpart_comparison.csv"
            ),
            "det_dense_compare": pd.read_csv(
                ROOT / "output/deterministic_dense_comparison.csv"
            ),
            "calibration": pd.read_csv(
                ROOT / "data/processed/calibration.csv"
            ),
        }
    )
    return {"config": config, **frames}


def f(value: float, digits: int = 1) -> str:
    return f"{value:,.{digits}f}"


def paragraph(text: str) -> dict:
    return {"kind": "paragraph", "text": text}


def equation(text: str) -> dict:
    return {"kind": "equation", "text": text}


def figure(key: str) -> dict:
    return {"kind": "figure", "key": key}


def table(key: str) -> dict:
    return {"kind": "table", "key": key}


def table_frame(inputs: dict, key: str) -> pd.DataFrame:
    return inputs[f"table_{TABLES[key][0]}"]


def number_objects(sections: list[dict]) -> tuple[list[dict], list[dict]]:
    assigned = {"fig": {}, "tab": {}}
    supplement_assigned = {"fig": {}, "tab": {}}
    placed = {"fig": 0, "tab": 0}
    supplement_items: list[dict] = []

    def replace(match: re.Match) -> str:
        kind, key = match.groups()
        registry = FIGURES if kind == "fig" else TABLES
        supplement = (
            SUPPLEMENT_FIGURES if kind == "fig" else SUPPLEMENT_TABLES
        )
        if key not in registry:
            raise KeyError(f"Unknown {kind} key {key}")
        if key in supplement:
            store = supplement_assigned[kind]
            if key not in store:
                store[key] = f"S{len(store) + 1}"
            label = "Figure" if kind == "fig" else "Table"
            return f"{label} {store[key]}"
        if key not in assigned[kind]:
            assigned[kind][key] = len(assigned[kind]) + 1
        label = "Figure" if kind == "fig" else "Table"
        return f"{label} {assigned[kind][key]}"

    for section in sections:
        kept_blocks = []
        for block in section["blocks"]:
            if block["kind"] == "paragraph":
                block["text"] = OBJECT_TOKEN.sub(replace, block["text"])
                kept_blocks.append(block)
            elif block["kind"] in {"figure", "table"}:
                kind = "fig" if block["kind"] == "figure" else "tab"
                key = block["key"]
                supplement = (
                    SUPPLEMENT_FIGURES if kind == "fig" else SUPPLEMENT_TABLES
                )
                if key in supplement:
                    if key not in supplement_assigned[kind]:
                        raise ValueError(
                            f"{kind}:{key} is placed before it is cited"
                        )
                    block["number"] = supplement_assigned[kind][key]
                    supplement_items.append(block)
                    continue
                if key not in assigned[kind]:
                    raise ValueError(f"{kind}:{key} is placed before it is cited")
                number = assigned[kind][key]
                if number != placed[kind] + 1:
                    raise ValueError(f"{kind}:{key} is placed out of order")
                placed[kind] = number
                block["number"] = number
                kept_blocks.append(block)
            else:
                kept_blocks.append(block)
        section["blocks"] = kept_blocks
    expected_main = {
        "fig": len(FIGURES) - len(SUPPLEMENT_FIGURES),
        "tab": len(TABLES) - len(SUPPLEMENT_TABLES),
    }
    for kind in ["fig", "tab"]:
        if placed[kind] != expected_main[kind]:
            raise ValueError(f"Not every main-text {kind} object is placed")
        registry = FIGURES if kind == "fig" else TABLES
        supplement = (
            SUPPLEMENT_FIGURES if kind == "fig" else SUPPLEMENT_TABLES
        )
        missing = supplement - set(supplement_assigned[kind])
        if missing:
            raise ValueError(f"Supplement {kind} objects never cited: {missing}")
        if set(registry) - supplement != set(assigned[kind]):
            raise ValueError(f"Main {kind} objects never cited")
    return sections, supplement_items


def build_revised_metrics(inputs: dict) -> dict:
    config = inputs["config"]
    size = inputs["size"]
    portfolio = inputs["portfolio"]
    baseline = size[size["regime"] == "baseline"]
    no_spoilage = size[size["regime"] == "full_carry_no_spoilage"]
    baseline_portfolios = portfolio[portfolio["regime"] == "baseline"]
    best_by_count = (
        baseline_portfolios.sort_values(["sku_count", "system_cost_jpy"])
        .groupby("sku_count")
        .first()
    )
    best_global = baseline_portfolios.loc[
        baseline_portfolios["system_cost_jpy"].idxmin()
    ]
    benchmark = baseline_portfolios[
        baseline_portfolios["is_current_market"].astype(str).str.lower().eq("true")
    ].iloc[0]
    thresholds = inputs["thresholds"][
        inputs["thresholds"]["regime"] == "baseline"
    ].set_index("from_sku_count")
    contrast = inputs["contrasts"][
        (inputs["contrasts"]["regime"] == "baseline")
        & (inputs["contrasts"]["contrast"] == "best_1_minus_best_2")
    ].iloc[0]
    consumption = inputs["consumption"].set_index(
        "annual_consumption_multiplier"
    )
    grid = inputs["grid"].set_index(["candidate_grid", "sku_count"])
    phase = inputs["phase"]
    phase_counts = phase["optimal_sku_count"].value_counts(normalize=True)
    portfolio_counts = phase["optimal_portfolio"].value_counts(normalize=True)
    robust = inputs["robust"].sort_values(
        ["maximum_regret_jpy", "mean_regret_jpy"]
    ).iloc[0]
    manufacturing = inputs["manufacturing"]
    high_sigma = manufacturing.loc[manufacturing["process_sigma_g"].idxmax()]
    parameter_counts = inputs["table_2"]["evidence_class"].value_counts()
    convergence = inputs["convergence"]
    return {
        "baseline_max_discard": baseline["physical_discard_kg"].max(),
        "no_spoilage_max_discard": no_spoilage["physical_discard_kg"].max(),
        "terminal_max_mean": inputs["terminal"]["mean_terminal_residual_g"].max(),
        "best_single": str(best_by_count.loc[1, "portfolio"]),
        "best_two": str(best_by_count.loc[2, "portfolio"]),
        "best_three": str(best_by_count.loc[3, "portfolio"]),
        "best_global": str(best_global["portfolio"]),
        "best_global_count": int(best_global["sku_count"]),
        "best_global_cost": float(best_global["system_cost_jpy"]),
        "one_to_two_threshold": float(
            thresholds.loc[
                1, "break_even_complexity_jpy_per_added_sku"
            ]
        ),
        "two_to_three_threshold": float(
            thresholds.loc[
                2, "break_even_complexity_jpy_per_added_sku"
            ]
        ),
        "one_to_two_nested": bool(
            thresholds.loc[1, "globally_optimal_portfolios_nested"]
        ),
        "two_vs_one_saving": float(contrast["mean_difference_jpy"]),
        "two_vs_one_ci_low": float(contrast["ci_low_jpy"]),
        "two_vs_one_ci_high": float(contrast["ci_high_jpy"]),
        "benchmark_regret": float(benchmark["regret_jpy"]),
        "consumption_low_single": str(
            consumption.loc[0.5, "best_one_sku_portfolio"]
        ),
        "consumption_base_single": str(
            consumption.loc[1.0, "best_one_sku_portfolio"]
        ),
        "consumption_high_single": str(
            consumption.loc[1.5, "best_one_sku_portfolio"]
        ),
        "consumption_low_two": str(
            consumption.loc[0.5, "best_two_sku_portfolio"]
        ),
        "consumption_high_two": str(
            consumption.loc[1.5, "best_two_sku_portfolio"]
        ),
        "consumption_min_kg": float(
            consumption["minimum_scenario_annual_consumption_kg"].min()
        ),
        "consumption_max_kg": float(
            consumption["maximum_scenario_annual_consumption_kg"].max()
        ),
        "coarse_single": str(grid.loc[("coarse", 1), "best_portfolio"]),
        "dense_single": str(grid.loc[("dense", 1), "best_portfolio"]),
        "coarse_two": str(grid.loc[("coarse", 2), "best_portfolio"]),
        "dense_two": str(grid.loc[("dense", 2), "best_portfolio"]),
        "coarse_three": str(grid.loc[("coarse", 3), "best_portfolio"]),
        "dense_three": str(grid.loc[("dense", 3), "best_portfolio"]),
        "dense_candidate_count": int(
            grid.loc[("dense", 1), "candidate_count"]
        ),
        "dense_min_kg": float(
            min(config["sensitivity"]["dense_candidate_package_kg"])
        ),
        "dense_max_kg": float(
            max(config["sensitivity"]["dense_candidate_package_kg"])
        ),
        "dense_step_kg": float(
            min(
                b - a
                for a, b in zip(
                    sorted(
                        set(
                            config["sensitivity"][
                                "dense_candidate_package_kg"
                            ]
                        )
                    ),
                    sorted(
                        set(
                            config["sensitivity"][
                                "dense_candidate_package_kg"
                            ]
                        )
                    )[1:],
                )
            )
        ),
        "national_annual_supply_kg": float(
            inputs["calibration"]
            .set_index("parameter")
            .loc[
                "national_annual_rice_supply_kg_per_person", "value"
            ]
        ),
        "phase_one_share": float(phase_counts.get(1, 0.0)),
        "phase_two_share": float(phase_counts.get(2, 0.0)),
        "phase_three_share": float(phase_counts.get(3, 0.0)),
        "phase_modal_portfolio": str(portfolio_counts.index[0]),
        "phase_modal_share": float(portfolio_counts.iloc[0]),
        "robust_portfolio": str(robust["portfolio"]),
        "robust_max_regret": float(robust["maximum_regret_jpy"]),
        "safety_margin": float(high_sigma["safety_margin_g"]),
        "high_sigma": float(high_sigma["process_sigma_g"]),
        "underfill_probability": float(
            high_sigma["maximum_underfill_probability"]
        ),
        "candidate_count": len(inputs["config"]["rice"]["candidate_package_kg"]),
        "household_count": len(inputs["config"]["households"]),
        "regime_count": len(inputs["config"]["regimes"]),
        "replications": int(inputs["config"]["project"]["replications"]),
        "sensitivity_replications": int(
            inputs["config"]["sensitivity"]["replications"]
        ),
        "scenario_complexity": float(
            inputs["config"]["costs"][
                "sku_complexity_jpy_per_household_year"
            ]
        ),
        "convergence_rate": float(convergence["converged"].mean()),
        "minimum_service_level": float(size["service_level_fraction"].min()),
        "unidentified_parameter_count": int(
            parameter_counts.get("unidentified", 0)
        ),
        "scenario_parameter_count": int(
            parameter_counts.get("scenario assumption", 0)
        ),
        "greedy_fail_count": int(
            (inputs["greedy"]["greedy_regret_at_2_jpy"] > 0).sum()
        ),
        "greedy_min_pair_regret": float(
            inputs["greedy"].loc[
                inputs["greedy"]["greedy_regret_at_2_jpy"] > 0,
                "greedy_regret_at_2_jpy",
            ].min()
        ),
        "greedy_max_pair_regret": float(
            inputs["greedy"]["greedy_regret_at_2_jpy"].max()
        ),
        "greedy_regime_count": int(len(inputs["greedy"])),
        "det_all_match": bool(
            inputs["det_compare"]["identities_match"].all()
        ),
        "det_decision_regret_max": float(
            inputs["det_compare"]["decision_regret_jpy"].max()
        ),
        "det_dense_match_count": int(
            inputs["det_dense_compare"]["identities_match"].sum()
        ),
        "det_dense_max_regret": float(
            inputs["det_dense_compare"]["decision_regret_jpy"].max()
        ),
        "det_dense_pick": dict(
            zip(
                inputs["det_dense_compare"]["sku_count"],
                inputs["det_dense_compare"][
                    "deterministic_optimal_portfolio"
                ],
            )
        ),
        "det_dense_stoch": dict(
            zip(
                inputs["det_dense_compare"]["sku_count"],
                inputs["det_dense_compare"][
                    "stochastic_optimal_portfolio"
                ],
            )
        ),
        "baseline_mcse_mean": float(
            inputs["convergence"]
            .query(
                "regime == 'baseline' and metric == 'generalized_cost_jpy'"
            )["mcse"]
            .mean()
        ),
        **_det_metrics(baseline, inputs["det_size"]),
    }


def _det_metrics(baseline: pd.DataFrame, det: pd.DataFrame) -> dict:
    baseline = baseline[baseline["regime"] == "baseline"]
    merged = det.merge(
        baseline[
            [
                "household",
                "package_kg",
                "generalized_cost_jpy",
                "storage_kg_days",
                "mean_age_at_consumption_days",
            ]
        ],
        on=["household", "package_kg"],
        suffixes=("_det", "_stoch"),
    )
    cost_delta = (
        merged["generalized_cost_jpy_det"] - merged["generalized_cost_jpy_stoch"]
    )
    cost_rel = (
        cost_delta.abs() / merged["generalized_cost_jpy_stoch"].abs().clip(lower=1e-9)
    )
    storage_rel = (
        (merged["storage_kg_days_det"] - merged["storage_kg_days_stoch"]).abs()
        / merged["storage_kg_days_stoch"].abs().clip(lower=1e-9)
    )
    age_rel = (
        (
            merged["mean_age_at_consumption_days_det"]
            - merged["mean_age_at_consumption_days_stoch"]
        ).abs()
        / merged["mean_age_at_consumption_days_stoch"].abs().clip(lower=1e-9)
    )
    return {
        "det_cost_bias_mean": float(cost_delta.mean()),
        "det_cost_max_rel": float(cost_rel.max()),
        "det_storage_max_rel": float(storage_rel.max()),
        "det_age_max_rel": float(age_rel.max()),
    }


def revised_abstract_text(metrics: dict) -> str:
    return (
        "Manufacturers must decide how many package sizes, and which capacities, to "
        "offer before heterogeneous consumption and household inventory carry-over are "
        "realized. This study couples a stochastic household carry-over inventory "
        "simulation—FIFO aging, spoilage, backlog, and final-batch adaptation—to "
        "exhaustive enumeration of package-SKU portfolios, and asks which parts of "
        "the decision actually need the stochastic layer. A globally evaluated "
        "SKU-inclusion rule follows: an added SKU is justified when its consumer "
        "value exceeds its complexity charge, without assuming nested optima; "
        "optimal portfolios are in fact non-nested, and a greedy add-one selection "
        f"lost {f(metrics['greedy_min_pair_regret'], 1)}\u2013"
        f"{f(metrics['greedy_max_pair_regret'], 1)} JPY/household-year at the "
        f"second SKU in {metrics['greedy_fail_count']} of "
        f"{metrics['greedy_regime_count']} regimes. A zero-discard boundary "
        "lemma, reproduced computationally, verifies that package–consumption "
        "mismatch alone cannot create physical discard without an inventory-exit "
        "mechanism. Matched discrete and continuous controls show the consumption "
        "lattice adds no discard and is a second-order cost term. A prespecified "
        "deterministic mean-trajectory ablation then shows which conclusions need "
        "the stochastic channels: on the canonical candidate set it recovers "
        "every optimal portfolio and household best size "
        f"(decision regret {f(metrics['det_decision_regret_max'], 1)} JPY), "
        f"while mechanism magnitudes do not transfer—storage differs by up to "
        f"{metrics['det_storage_max_rel']:.0%}, age by up to "
        f"{metrics['det_age_max_rel']:.0%}, and spoilage windows, adaptation, "
        "and backlog exist only stochastically. "
        "In transparent rice scenarios, the one-to-two and two-to-three SKU "
        f"break-even charges were {f(metrics['one_to_two_threshold'], 1)} and "
        f"{f(metrics['two_to_three_threshold'], 1)} JPY/household-year; under "
        f"the assumed {metrics['scenario_complexity']:.0f} JPY charge, "
        f"{metrics['best_two']} kg was optimal. Dense phase maps separate one-, "
        "two-, and three-SKU regions; exact portfolios remain "
        "calibration-dependent rather than market recommendations."
    )


def build_revised_sections(inputs: dict, metrics: dict) -> list[dict]:
    config = inputs["config"]
    sensitivity = config["sensitivity"]
    household_choices = ", ".join(
        f"{row.household}: {row.package_kg:g} kg"
        for row in (
            inputs["size"][inputs["size"]["regime"] == "baseline"]
            .sort_values("generalized_cost_jpy")
            .groupby("household", as_index=False)
            .first()
            .itertuples(index=False)
        )
    )
    households = table_frame(inputs, "households")
    spacing_text = "; ".join(
        f"{spacing:g} g for {', '.join(group['scenario'])}"
        for spacing, group in households.groupby("lattice_spacing_g")
    )
    factorial = table_frame(inputs, "factorial")
    factorial_best = factorial[
        factorial["package_kg"] == float(metrics["best_single"])
    ].iloc[0]
    max_lattice_discard = factorial["discard_lattice_kg"].abs().max()
    convergence = inputs["convergence"]
    relative_threshold = float(config["project"]["relative_mcse_threshold"])
    transition_labels = {
        "best_one_sku_portfolio": "the best single size",
        "best_two_sku_portfolio": "the best two-SKU portfolio",
        "best_three_sku_portfolio": "the best three-SKU portfolio",
        "global_optimal_portfolio": "the global optimum",
    }
    transition_text = "; ".join(
        f"{transition_labels[row.outcome]} changed from {row.from_portfolio} to "
        f"{row.to_portfolio} kg between multipliers {row.lower_multiplier:g} and "
        f"{row.upper_multiplier:g}"
        for row in inputs["consumption_transitions"].itertuples(index=False)
    )
    consumption_table = table_frame(inputs, "consumption")
    phase_thresholds = table_frame(inputs, "phase_thresholds")
    phase_pairs = len(phase_thresholds)
    regret = table_frame(inputs, "regret")
    robust_regime = (
        str(regret.iloc[0]["regime_of_maximum"])
        .replace("_", " ")
        .replace("sku", "SKU")
    )
    manufacturing = table_frame(inputs, "manufacturing")
    largest_giveaway = manufacturing.loc[
        manufacturing["expected_giveaway_percent"].idxmax()
    ]
    targets = ", ".join(f"{value:g}" for value in manufacturing["target_kg"].unique())
    multiplier_step = (
        sensitivity["annual_consumption_multipliers"][1]
        - sensitivity["annual_consumption_multipliers"][0]
    )
    return [
        {
            "heading": "1. Introduction",
            "blocks": [
                paragraph(
                    "Choosing how many package sizes to offer, and which capacities, "
                    "is a recurring industrial-engineering decision linking nominal "
                    "capacity, purchasing frequency, inventory exposure, packaging "
                    "material, handling, and assortment complexity. Prior work optimizes "
                    "package-size sets at the supply-chain level—outer-pack "
                    "combinations, p-median size selection, and retail case-pack "
                    "sizes—or prices a single size under usage heterogeneity, but "
                    "does not represent the downstream household inventory dynamics that "
                    "give a consumer-size SKU its value (Gámez Albán et al., 2015; "
                    "Brinker & Gündüz, 2016; Wensing et al., 2018; Zhao et al., 2023; "
                    "Koenigsberg et al., 2010)."
                ),
                paragraph(
                    "The decision also belongs to the broader product-variety and "
                    "assortment literature. Product-line and pack-size breadth can serve "
                    "heterogeneous demand, but variety also changes inventory, assembly, "
                    "supply-chain, forecasting, setup, and planning costs (Kekre & "
                    "Srinivasan, 1990; Fisher & Ittner, 1999; van Ryzin & Mahajan, 1999; "
                    "Randall & Ulrich, 2001; Wan et al., 2014; Wan & Sanders, 2017; "
                    "Malinowski et al., 2018; Menezes et al., 2021). That literature "
                    "establishes the value–complexity trade-off; it does not supply the "
                    "inventory-exit boundary or the globally evaluated SKU-inclusion rule "
                    "for package capacities developed here. Packaging and food-waste "
                    "research likewise shows that package size can affect consumer "
                    "practice, food loss, and packaging burdens, with effects that depend "
                    "on the modeled mechanism and system boundary (Chan, 2022; Fresán et "
                    "al., 2019; Pauer et al., 2019; van Rooijen et al., 2024; Williams et "
                    "al., 2020; Wilson et al., 2017). Household-level simulation and "
                    "planning models evaluate waste under given retail offerings but do "
                    "not choose the offered package-size set (Kandemir et al., 2020; van "
                    "Rooijen et al., 2024)."
                ),
                paragraph(
                    "A recurring shortcut is to compute a remainder when package content "
                    "is divided by a nominal consumption unit and interpret that remainder "
                    "as waste. The arithmetic identifies a possible inventory state, not "
                    "a physical exit. Consumers can carry inventory, combine packages, "
                    "adapt a final batch, postpone demand, or dispose of product. These "
                    "responses have different material-flow and service consequences."
                ),
                paragraph(
                    "The unresolved production problem has two levels. First, when does "
                    "consumption granularity materially change package-level cost once "
                    "carry-over, aging, and purchasing are represented? Second, when does "
                    "the consumer value of an additional package-size SKU exceed the "
                    "producer's manufacturing, forecasting, stocking, and assortment "
                    "complexity? A package optimum without a SKU-inclusion rule cannot "
                    "answer the second question."
                ),
                paragraph(
                    "Simulation has become the working instrument for household "
                    "inventory and food-waste questions: discrete-event household "
                    "models evaluate waste interventions under fixed retail "
                    "offerings (Kandemir et al., 2020; Boz et al., 2026), "
                    "dynamic stochastic household models embed lumpy package "
                    "purchases in inventory processes (Zhang & Qi, 2020), and "
                    "perishable-inventory applications in this journal couple "
                    "stochastic simulation to optimization loops on the supply "
                    "side (Onggo et al., 2019; Moubed et al., 2021). What is "
                    "missing is the demand-side counterpart: a household "
                    "inventory simulation that generates each candidate SKU's "
                    "value endogenously and closes into the producer's "
                    "portfolio-selection problem, together with a transparent "
                    "answer to the question every such coupling raises—does the "
                    "stochastic layer actually change the decision?"
                ),
                paragraph(
                    "This study contributes (i) an integrated decision model coupling "
                    "stochastic household carry-over inventory to package-SKU portfolio "
                    "choice; (ii) a zero-discard boundary lemma separating lattice "
                    "mismatch from inventory exit, used as an internal verification "
                    "anchor the simulator must reproduce; (iii) a global SKU-inclusion "
                    "condition whose necessity is demonstrated, not assumed: optimal "
                    "portfolios are non-nested, and nested greedy selection loses "
                    f"{f(metrics['greedy_min_pair_regret'], 1)}\u2013"
                    f"{f(metrics['greedy_max_pair_regret'], 1)} JPY/household-year "
                    f"at the second SKU in {metrics['greedy_fail_count']} of "
                    f"{metrics['greedy_regime_count']} prespecified regimes; "
                    "(iv) matched discrete, continuous, and no-spoilage negative "
                    "controls; (v) a deterministic mean-trajectory ablation that "
                    "localizes exactly which conclusions require the stochastic "
                    "channels—portfolio and household-level optima are recovered "
                    "deterministically, while storage, age, spoilage-window, "
                    "adaptation, and backlog magnitudes are not; and (vi) dense "
                    "phase maps with annual-consumption and "
                    "candidate-grid robustness analyses. Rice supplies a transparent "
                    "application and public calibration anchors, not the intellectual "
                    "premise or a claimed population estimate."
                ),
            ],
        },
        {
            "heading": "2. Model and methods",
            "blocks": [
                {
                    "kind": "subheading",
                    "text": "2.1. Inventory states and estimands",
                },
                paragraph(
                    "<<tab:definitions>> defines purchases, demand, served consumption, "
                    "backlog, carried inventory, and inventory exit. <<fig:concept>> "
                    "separates residual inventory from physical discard. Generalized cost "
                    "is reported separately from physical quantities so that monetary "
                    "assumptions do not redefine material waste."
                ),
                table("definitions"),
                figure("concept"),
                {
                    "kind": "subheading",
                    "text": "2.2. Zero-discard boundary (Lemma 1)",
                },
                paragraph(
                    "For periods t=1,…,T, let Iₜ be usable inventory after the event, "
                    "Qₜ purchased quantity, Dₜ demand, Cₜ served consumption, Bₜ "
                    "backlog, and Eₜ physical inventory exit, with nonnegative initial "
                    "states I₀ and B₀. The flow equations are:"
                ),
                equation("Bₜ = Bₜ₋₁ + Dₜ − Cₜ,"),
                equation("Iₜ = Iₜ₋₁ + Qₜ − Cₜ − Eₜ."),
                paragraph(
                    "Lemma 1. If all unconsumed usable inventory is carried beyond "
                    "each event and the reporting horizon; no spoilage, expiry, forced "
                    "removal, leakage, damage, capacity overflow, or terminal disposal "
                    "occurs; consumption cannot exceed usable inventory; and unserved "
                    "demand remains backlog, then total physical discard is zero for "
                    "every package size, demand path, and package–consumption remainder."
                ),
                paragraph(
                    "Proof (consistency check). The assumptions imply Eₜ=0 for every "
                    "t. Summing the inventory "
                    "equation gives I_T−I₀=Σₜ(Qₜ−Cₜ−Eₜ), and therefore "
                    "I_T=I₀+ΣₜQₜ−ΣₜCₜ≥0 when Eₜ=0. A package lattice can change "
                    "purchase timing, purchase count, and terminal inventory, but it "
                    "cannot introduce a positive Eₜ. Summing Eₜ=0 over the horizon gives "
                    "ΣₜEₜ=0. If service is infeasible, the backlog balance places the "
                    "unserved quantity in B_T rather than discard. The result fails only "
                    "when an explicit inventory-exit mechanism or terminal disposal rule "
                    "is introduced."
                ),
                paragraph(
                    "Stockouts therefore enter backlog and service loss, not discard. "
                    "The conclusion changes when terminal inventory is disposed, shelf "
                    "life is finite, carry-over is prohibited, product is abandoned, or "
                    "another inventory-exit rule is imposed. Initial inventory and "
                    "non-discard transfers must then be retained explicitly in the "
                    "balance."
                ),
                {
                    "kind": "subheading",
                    "text": "2.3. Stochastic inventory and cost model",
                },
                paragraph(
                    "Consumption events follow a gamma renewal process. Event quantities "
                    "are drawn from scenario-specific discrete lattices; matched "
                    "lognormal and gamma controls preserve each scenario's event-quantity "
                    "mean and variance while removing the lattice. Inventory is consumed "
                    "FIFO. After a freshness threshold, lot quantity follows continuous "
                    "exponential decay; storage exposure integrates the same decay path. "
                    "Freshness cost is charged lot by lot, preventing fresh inventory "
                    "from diluting the age exposure of older consumed lots."
                ),
                paragraph(
                    "If remaining inventory is close to planned usage, adaptation can "
                    "defer the unserved quantity as backlog. Otherwise enough packages "
                    "are purchased to serve planned usage. The objective adds product, "
                    "unmet-demand, purchase-event, packaging, storage, freshness, "
                    "handling, and adaptation costs. Physical discard, packaging mass, "
                    "age, purchases, service, and terminal inventory remain separate "
                    "outputs. Product cost is a real purchase outlay at effective "
                    "prices; because total consumption is fixed by demand, cost "
                    "differences across package sizes arise from spoilage-driven "
                    "repurchasing, purchase frequency, and the size-discount term "
                    "rather than from the transfer level of spending itself."
                ),
                paragraph(
                    "Public sources anchor the 150 g rice unit, a 30-day storage guide, "
                    f"{metrics['national_annual_supply_kg']:g} kg annual "
                    "per-person supply, household composition, and observed "
                    "2, 5, and 10 kg package examples (Ministry of Agriculture, "
                    "Forestry and Fisheries of Japan [MAFF], 2026a, 2026b, 2026c; Mino "
                    "Rice Co., Ltd., 2026; Statistics Bureau of Japan, 2025). <<tab:parameters>> "
                    "classifies every important input as empirically estimated, "
                    "empirically anchored, a scenario assumption, or unidentified. The "
                    "complete machine-readable evidence table supplied with the "
                    "reproducibility package retains source identifiers and limitations."
                ),
                table("parameters"),
                paragraph(
                    f"{metrics['household_count']} household design scenarios span "
                    "one-person, two-person, and three- and four-person family "
                    "consumption (<<tab:households>>). Annual consumption ranges from "
                    f"{f(households['annual_consumption_kg'].min(), 1)} to "
                    f"{f(households['annual_consumption_kg'].max(), 1)} kg, and each "
                    "scenario has its own discrete usage quanta and probabilities, "
                    "final-batch adaptation probability and minimum ratio, freshness "
                    "threshold, and handling threshold. The weights preserve only the "
                    "census split between one-person and other households; the classes "
                    "are transparent stress-test scenarios, not fitted demographic "
                    "segments."
                ),
                table("households"),
                {
                    "kind": "subheading",
                    "text": "2.4. Package-portfolio objective and the SKU-inclusion rule",
                },
                paragraph(
                    "Product-variety research documents both demand-side benefits and "
                    "inventory, forecasting, manufacturing, and network costs (Wan et "
                    "al., 2014; Wan & Sanders, 2017; Malinowski et al., 2018; Menezes et "
                    "al., 2021). "
                    "Here those producer-side burdens are represented transparently by "
                    "a stated complexity function rather than treated as estimated costs. "
                    "For offered set K, scenario h with weight w_h, and package cost "
                    "C_h(P), define:"
                ),
                equation("J(K) = Σ_h w_h min_{P∈K} C_h(P) + C_complexity(K)."),
                paragraph(
                    "Let V_k be the minimum consumer cost over all globally enumerated "
                    "k-SKU portfolios. With linear complexity c per SKU beyond the first, "
                    "J_k=V_k+(k−1)c. The SKU-inclusion rule follows directly: moving from the "
                    "global k-SKU optimum to the global (k+1)-SKU optimum improves cost "
                    "if and only if c<V_k−V_{k+1}; equality gives a tie. The portfolios "
                    "need not be nested, and the global SKU count minimizes J_k across "
                    "all k. This algebraic threshold applies to the stated linear "
                    "complexity model; nonlinear or portfolio-specific producer costs "
                    "require direct evaluation of the corresponding C_complexity(K)."
                ),
                {
                    "kind": "subheading",
                    "text": "2.5. Computational design and uncertainty",
                },
                paragraph(
                    f"The canonical experiment evaluates {metrics['candidate_count']} "
                    f"sizes, {metrics['household_count']} scenarios, and "
                    f"{metrics['regime_count']} regimes with "
                    f"{metrics['replications']} replications per cell after a one-year "
                    "burn-in. Common random numbers align package sizes. Replication-level "
                    "scenario aggregation preserves covariance in Monte Carlo standard "
                    "errors. Sensitivity runs use "
                    f"{metrics['sensitivity_replications']} replications per cell and "
                    "retain common random numbers within each comparison."
                ),
                paragraph(
                    f"<<tab:regimes>> lists the {metrics['regime_count']} boundary and "
                    "sensitivity regimes. Each removes or amplifies a distinct "
                    "mechanism—spoilage, final-batch adaptation, SKU complexity, "
                    "packaging and purchase logistics, consumer inconvenience, or the "
                    "discrete usage lattice—so that a result generated by aging, "
                    "adaptation, logistics, or complexity is not attributed to discrete "
                    "mismatch."
                ),
                table("regimes"),
                paragraph(
                    "The primary negative control crosses discrete versus matched "
                    "continuous usage with spoilage on versus off. Dense phase analysis "
                    "varies spoilage intensity and purchase-event cost over a "
                    f"{len(sensitivity['spoilage_hazard_multipliers'])}×"
                    f"{len(sensitivity['purchase_event_cost_multipliers'])} grid of "
                    "multipliers and evaluates linear SKU complexity from "
                    f"{min(sensitivity['sku_complexity_jpy_per_household_year']):g} to "
                    f"{max(sensitivity['sku_complexity_jpy_per_household_year']):g} "
                    "JPY/household-year per added SKU at every pair. Annual consumption "
                    "scales all six scenarios together from "
                    f"{min(sensitivity['annual_consumption_multipliers']):g} to "
                    f"{max(sensitivity['annual_consumption_multipliers']):g} of their "
                    f"anchors in steps of {multiplier_step:.1f}, testing whether package "
                    "and portfolio identities depend on the original consumption values. "
                    f"Candidate-grid robustness adds {metrics['dense_step_kg']:g} kg "
                    f"increments from {metrics['dense_min_kg']:g} to "
                    f"{metrics['dense_max_kg']:g} kg "
                    "to the original anchors and exhaustively enumerates all one-, two-, "
                    f"and three-SKU portfolios over {metrics['dense_candidate_count']} "
                    "sizes."
                ),
                paragraph(
                    "Monte Carlo convergence was assessed for every regime–scenario–size "
                    "cell and four reported metrics: generalized cost, physical discard, "
                    "packaging mass, and age at consumption. A diagnostic passed when "
                    f"relative MCSE was at most {relative_threshold:.0%}, when MCSE was "
                    "below a metric-specific absolute tolerance, or when mean and MCSE "
                    f"were both exactly zero; the rule passed "
                    f"{metrics['convergence_rate']:.1%} of {len(convergence):,} "
                    "diagnostics. Absolute tolerances prevent near-zero means, such as "
                    "discard without spoilage, from failing solely because relative MCSE "
                    "is unstable. Monte Carlo intervals quantify simulation error "
                    "conditional on parameters; they do not quantify parameter "
                    "uncertainty. The pipeline uses frozen public-source snapshots with "
                    "SHA-256 validation, centralized YAML configuration, deterministic "
                    "seed derivation, and code-generated results; raw files are never "
                    "generated from manuscript values."
                ),
                {
                    "kind": "subheading",
                    "text": "2.6. Deterministic counterpart (ablation design)",
                },
                paragraph(
                    "To localize what the stochastic layer contributes, the same "
                    "household–package grid and portfolio objective are evaluated "
                    "under a deterministic mean-trajectory counterpart: the "
                    "inter-purchase interval is fixed at its mean, each event's "
                    "usage equals its expected quantity, and adaptive consumption "
                    "is removed so purchases occur deterministically when needed. "
                    "This is the spreadsheet-style accounting a deterministic "
                    "analysis would produce. The prespecified comparison asks "
                    "(i) whether the deterministic counterpart recovers the same "
                    "optimal portfolio at every SKU count and the same "
                    "household-level best sizes, and (ii) the decision regret of "
                    "the deterministically chosen portfolio re-evaluated under "
                    "the stochastic model, with mechanism-level differences "
                    "(storage exposure, age at consumption, discard, backlog, "
                    "adaptation) reported separately."
                ),
            ],
        },
        {
            "heading": "3. Results",
            "blocks": [
                {
                    "kind": "subheading",
                    "text": "3.1. Analytical zero-discard boundary",
                },
                paragraph(
                    "Lemma 1 establishes the central boundary: modulo mismatch can "
                    "change terminal inventory but cannot create physical discard without "
                    "an exit mechanism. Lattice compatibility is diagnosed by each "
                    "scenario's integer lattice spacing, the greatest common divisor of "
                    f"its usage quanta ({spacing_text}; <<tab:households>>). "
                    "Compatibility is necessary for exact depletion on the lattice but "
                    "does not imply zero sequence-specific residual. The deterministic "
                    "one-package residual P mod 150 g is periodic in nominal size "
                    "(<<fig:lattice>>), and the largest mean one-package terminal "
                    f"residual was {f(metrics['terminal_max_mean'], 1)} g; neither "
                    "diagnostic quantity is a waste estimate."
                ),
                figure("lattice"),
                {
                    "kind": "subheading",
                    "text": "3.2. Computational validation",
                },
                paragraph(
                    "The full-carry-over/no-spoilage regime produced a maximum of "
                    f"{f(metrics['no_spoilage_max_discard'], 3)} "
                    "kg/household-year physical discard across all scenario-size cells. "
                    "The matched continuous no-spoilage control also remained at zero. "
                    "These simulations verify implementation consistency with Lemma "
                    "1; they are not empirical proof of household behavior. "
                    "<<fig:controls>> and <<tab:controls>> report the matched controls."
                ),
                figure("controls"),
                table("controls"),
                {
                    "kind": "subheading",
                    "text": "3.3. Mechanisms producing economic consequences",
                },
                paragraph(
                    "With aging and spoilage active, maximum cell-level discard reached "
                    f"{f(metrics['baseline_max_discard'], 3)} kg/household-year. "
                    "Removing the lattice did not eliminate discard because continuous "
                    "usage still left inventory exposed to aging. Conversely, removing "
                    "spoilage eliminated discard despite discrete usage. Cost differences "
                    "therefore operate through purchase frequency, product loss, "
                    "storage, freshness, packaging, handling, and their interaction with "
                    "spoilage rather than through a remainder interpreted as waste."
                ),
                paragraph(
                    "<<tab:factorial>> decomposes each package size into the lattice "
                    "main effect under no spoilage, the spoilage main effect under "
                    "matched continuous usage, and their interaction. At the baseline "
                    f"best single size of {metrics['best_single']} kg, these components "
                    "of generalized cost were "
                    f"{f(factorial_best['cost_lattice_jpy'], 1)}, "
                    f"{f(factorial_best['cost_spoilage_jpy'], 1)}, and "
                    f"{f(factorial_best['cost_interaction_jpy'], 1)} "
                    "JPY/household-year, respectively. The largest absolute lattice main "
                    "effect on physical discard across package sizes was "
                    f"{f(max_lattice_discard, 3)} kg/household-year, consistent with "
                    "Lemma 1. Component differences are accounting pathways rather "
                    "than independent causal effects."
                ),
                table("factorial"),
                {
                    "kind": "subheading",
                    "text": "3.4. Consumption heterogeneity and single-size choice",
                },
                paragraph(
                    "<<tab:package_comparison>> reports the scenario-weighted baseline "
                    f"comparison of all {metrics['candidate_count']} candidate sizes, "
                    "including Monte Carlo intervals for generalized cost and discard, "
                    "purchases, packaging mass, age at consumption, and regret relative "
                    "to household-specific optima. At the anchor, household-specific "
                    f"single-size minima were {household_choices}. These are "
                    "design-scenario optima, not estimated recommendations for named "
                    "household populations."
                ),
                table("package_comparison"),
                paragraph(
                    "Across the annual-consumption scale, scenario values ranged from "
                    f"{f(metrics['consumption_min_kg'], 1)} to "
                    f"{f(metrics['consumption_max_kg'], 1)} kg/year. The optimal single "
                    f"size changed from {metrics['consumption_low_single']} kg at the low "
                    f"scale to {metrics['consumption_high_single']} kg at the high scale; "
                    f"the anchor retained {metrics['consumption_base_single']} kg. "
                    f"Specifically, {transition_text}. Regret of the 2|5|10 kg benchmark "
                    f"ranged from {f(consumption_table['benchmark_regret_jpy'].min(), 1)} "
                    f"to {f(consumption_table['benchmark_regret_jpy'].max(), 1)} "
                    "JPY/household-year. Accordingly, the baseline single-size result is "
                    "not invariant to throughput (<<fig:consumption>> and "
                    "<<tab:consumption>>)."
                ),
                figure("consumption"),
                table("consumption"),
                {
                    "kind": "subheading",
                    "text": "3.5. Portfolio value of variety",
                },
                paragraph(
                    f"The globally best one-, two-, and three-SKU portfolios before "
                    f"additional complexity were {metrics['best_single']}, "
                    f"{metrics['best_two']}, and {metrics['best_three']} kg. At the "
                    f"scenario complexity charge of {metrics['scenario_complexity']:.0f} "
                    f"JPY/household-year per added SKU and the coarse candidate set, "
                    f"{metrics['best_global']} kg with "
                    f"{metrics['best_global_count']} SKUs minimized system cost. The paired "
                    f"single-minus-two-SKU difference was "
                    f"{f(metrics['two_vs_one_saving'], 1)} JPY/household-year "
                    f"(95% Monte Carlo interval "
                    f"{f(metrics['two_vs_one_ci_low'], 1)}–"
                    f"{f(metrics['two_vs_one_ci_high'], 1)}). <<fig:pareto>> shows the "
                    "cost–packaging frontier."
                ),
                figure("pareto"),
                {
                    "kind": "subheading",
                    "text": "3.6. SKU break-even thresholds",
                },
                paragraph(
                    "The scenario one-to-two break-even complexity charge was "
                    f"{f(metrics['one_to_two_threshold'], 1)} "
                    "JPY/household-year per added SKU; the two-to-three threshold was "
                    f"{f(metrics['two_to_three_threshold'], 1)}. Thus a second SKU pays "
                    "over a much wider complexity range than a third. The global one- and "
                    "two-SKU optima were not nested, so an add-one heuristic would not "
                    "recover the correct comparison (<<fig:break_even>> and "
                    "<<tab:thresholds>>)."
                ),
                paragraph(
                    "This non-nestedness is material, not decorative. <<tab:greedy>> "
                    "compares a nested greedy construction\u2014retain the best "
                    "(k\u22121)-SKU set and add the argmin next SKU\u2014against the "
                    f"global optimum in all {metrics['greedy_regime_count']} "
                    f"prespecified regimes. Greedy missed the optimal two-SKU "
                    f"portfolio in {metrics['greedy_fail_count']} regimes, always "
                    "by locking in the best single size, and the loss reached "
                    f"{f(metrics['greedy_min_pair_regret'], 1)}\u2013"
                    f"{f(metrics['greedy_max_pair_regret'], 1)} JPY/household-year. "
                    "The failure is confined to the first inclusion decision: once "
                    "the greedy two-SKU set was fixed, the added third SKU matched "
                    "the global optimum in every regime. The marginal inclusion "
                    "rule must therefore be applied to globally evaluated optima, "
                    "not to incrementally built sets."
                ),
                figure("break_even"),
                table("thresholds"),
                table("greedy"),
                {
                    "kind": "subheading",
                    "text": "3.7. Dense phase and regime analysis",
                },
                paragraph(
                    "Across the dense spoilage–purchase–complexity phase map, one-, two-, "
                    f"and three-SKU systems were optimal in "
                    f"{metrics['phase_one_share']:.1%}, "
                    f"{metrics['phase_two_share']:.1%}, and "
                    f"{metrics['phase_three_share']:.1%} of cells, respectively. "
                    f"{metrics['phase_modal_portfolio']} kg was the modal portfolio "
                    f"({metrics['phase_modal_share']:.1%} of cells). Higher complexity "
                    "contracts variety; stronger purchase burden tends to favor larger "
                    "sizes, while stronger aging loss can increase the value of smaller "
                    "alternatives (<<fig:phase>>)."
                ),
                figure("phase"),
                paragraph(
                    "<<tab:phase_thresholds>> reports the global thresholds for each of "
                    f"the {phase_pairs} spoilage–purchase pairs. One-to-two thresholds "
                    "ranged from "
                    f"{f(phase_thresholds['one_to_two_threshold_jpy'].min(), 1)} to "
                    f"{f(phase_thresholds['one_to_two_threshold_jpy'].max(), 1)} and "
                    "two-to-three thresholds from "
                    f"{f(phase_thresholds['two_to_three_threshold_jpy'].min(), 1)} to "
                    f"{f(phase_thresholds['two_to_three_threshold_jpy'].max(), 1)} "
                    "JPY/household-year per added SKU; the globally optimal one- and "
                    "two-SKU portfolios were nested in "
                    f"{int(phase_thresholds['one_two_nested'].sum())} of {phase_pairs} "
                    "pairs. Phase-map shares and thresholds are scenario-map quantities, "
                    "not population probabilities or observed producer costs."
                ),
                table("phase_thresholds"),
                {
                    "kind": "subheading",
                    "text": "3.8. Candidate-grid robustness",
                },
                paragraph(
                    f"The coarse-grid one-, two-, and three-SKU optima were "
                    f"{metrics['coarse_single']}, {metrics['coarse_two']}, and "
                    f"{metrics['coarse_three']} kg. Exhaustive dense-grid enumeration "
                    f"returned {metrics['dense_single']}, {metrics['dense_two']}, and "
                    f"{metrics['dense_three']} kg. The one-SKU identity and broad "
                    "small-plus-large portfolio structure persisted, but the exact "
                    "two- and three-SKU identities changed. The coarse-grid "
                    f"{metrics['coarse_two']} kg result is therefore conditional on the "
                    f"{metrics['scenario_complexity']:.0f} JPY/household-year complexity "
                    "charge and candidate grid, not a uniquely robust physical optimum "
                    "or a Japanese-market recommendation (<<fig:grid>> and "
                    "<<tab:grid>>)."
                ),
                figure("grid"),
                table("grid"),
                {
                    "kind": "subheading",
                    "text": "3.9. Robust portfolio and benchmark regret",
                },
                paragraph(
                    "The verified 2|5|10 kg manufacturer example had anchor-regime regret "
                    f"of {f(metrics['benchmark_regret'], 1)} JPY/household-year. The "
                    f"prespecified-regime minimax-regret portfolio was "
                    f"{metrics['robust_portfolio']} kg with maximum regret "
                    f"{f(metrics['robust_max_regret'], 1)} JPY/household-year, arising "
                    f"in the {robust_regime} regime (<<tab:regret>>). Scenario-set "
                    "regret is the maximum within-regime regret over the declared "
                    "regimes; it is robustness over a declared scenario set, not a "
                    "probability model for parameter uncertainty."
                ),
                table("regret"),
                {
                    "kind": "subheading",
                    "text": "3.10. Fill-process precision",
                },
                paragraph(
                    "Nominal package size, fill-process overage, and package–consumption "
                    "mismatch are distinct. Under a normal fill process with nominal "
                    "target T, process standard deviation σ, and maximum underfill "
                    "probability α, the required mean fill is:"
                ),
                equation("μ = T + z_(1−α) σ."),
                paragraph(
                    f"Across targets of {targets} kg and σ from "
                    f"{manufacturing['process_sigma_g'].min():g} to "
                    f"{manufacturing['process_sigma_g'].max():g} g, the safety margin "
                    f"rose to {f(metrics['safety_margin'], 1)} g at σ="
                    f"{metrics['high_sigma']:.0f} g and α="
                    f"{metrics['underfill_probability']:.3f} (<<tab:manufacturing>>). "
                    "Because the margin is absolute, the expected relative giveaway was "
                    "largest for the smallest target, reaching "
                    f"{f(largest_giveaway['expected_giveaway_percent'], 2)}% at "
                    f"{largest_giveaway['target_kg']:g} kg and σ="
                    f"{largest_giveaway['process_sigma_g']:g} g (<<fig:manufacturing>>). "
                    "The α value is an analytical scenario constraint, not a legal "
                    "standard."
                ),
                table("manufacturing"),
                figure("manufacturing"),
                {
                    "kind": "subheading",
                    "text": "3.11. Deterministic counterpart ablation",
                },
                paragraph(
                    "<<tab:deterministic>> compares portfolio selection under the "
                    "deterministic mean-trajectory counterpart and the stochastic "
                    "model. The deterministic counterpart recovers the identical "
                    "optimal portfolio at every SKU count "
                    f"({metrics['best_single']} | {metrics['best_two']} | "
                    f"{metrics['best_three']} kg) and the identical household-level "
                    "best size for all six household scenarios; re-evaluating the "
                    "deterministically chosen portfolios under the stochastic model "
                    "gives zero decision regret. The stochastic channels therefore "
                    "do not change which portfolio is optimal in the studied "
                    "calibration—a boundary worth reporting rather than hiding. "
                    "On the dense 39-size candidate grid the agreement degrades only "
                    "to adjacent-size swaps (deterministic picks "
                    f"{metrics['det_dense_pick'][1]}, "
                    f"{metrics['det_dense_pick'][2]}, and "
                    f"{metrics['det_dense_pick'][3]} kg versus the stochastic "
                    f"{metrics['det_dense_stoch'][1]}, "
                    f"{metrics['det_dense_stoch'][2]}, and "
                    f"{metrics['det_dense_stoch'][3]} kg; identities match in "
                    f"{metrics['det_dense_match_count']} of 3 SKU counts and the "
                    f"largest decision regret is "
                    f"{f(metrics['det_dense_max_regret'], 1)} JPY/household-year, "
                    f"below the {f(metrics['baseline_mcse_mean'], 0)} JPY mean "
                    "per-cell Monte-Carlo error). Identity-level agreement holds on "
                    "the canonical candidate set; near-ties between adjacent sizes "
                    "on finer grids are below stochastic resolution."
                ),
                paragraph(
                    "The channels do change magnitudes. The deterministic model "
                    "deviates from the stochastic model's generalized cost by an "
                    f"average of {f(metrics['det_cost_bias_mean'], 1)} "
                    "JPY/household-year (maximum relative deviation "
                    f"{metrics['det_cost_max_rel']:.1%}), storage exposure differs "
                    f"by up to {metrics['det_storage_max_rel']:.0%}, and age at "
                    f"consumption by up to {metrics['det_age_max_rel']:.0%}. "
                    "Adaptive consumption events, spoilage-window discard, and "
                    "backlog episodes exist only under the stochastic channels, so "
                    "any conclusion about mechanism magnitudes, service, or waste "
                    "pathways—not merely the portfolio argmin—requires the "
                    "simulation. The non-nested optimal portfolios also appear "
                    "deterministically, identifying household heterogeneity, not "
                    "randomness, as their source."
                ),
                table("deterministic"),
            ],
        },
        {
            "heading": "4. Discussion",
            "blocks": [
                {
                    "kind": "subheading",
                    "text": "4.1. Structural and scenario-dependent results",
                },
                paragraph(
                    "Three results are structural. Residual inventory is not physical "
                    "waste; discard requires an inventory-exit mechanism; and an added SKU "
                    "is justified when its globally evaluated consumer benefit exceeds "
                    "its complexity charge. Package identities, monetary savings, discard "
                    "levels, numerical thresholds, phase-map proportions, and minimax "
                    "values are scenario-dependent because behavior and producer cost "
                    "were not estimated. The dense grid preserves broad portfolio "
                    "structure while changing exact multi-SKU identities."
                ),
                paragraph(
                    "Relative to the closest prior work, the advance is integrative "
                    "rather than algorithmic. Size-set optimizers choose packages under "
                    "deterministic demand without household dynamics (Gámez Albán et "
                    "al., 2015; Brinker & Gündüz, 2016; Wensing et al., 2018; Zhao et "
                    "al., 2023); single-size design models treat usage heterogeneity "
                    "without inventory states (Koenigsberg et al., 2010); and household "
                    "simulation or planning models evaluate waste under a fixed retail "
                    "offering (Kandemir et al., 2020; van Rooijen et al., 2024). The "
                    "present model is the demand-side complement to case-pack "
                    "optimization: it prices each candidate consumer SKU inside a "
                    "stochastic household inventory process and then selects the offered "
                    "portfolio under an explicit complexity charge. Assortment-plus-"
                    "inventory models assume an exogenous choice or substitution value "
                    "function (Gaur & Honhon, 2006; van Ryzin & Mahajan, 1999; K\u00f6k "
                    "& Fisher, 2007); here each SKU's value is instead generated "
                    "endogenously by the inventory process. The quantified consequence "
                    "is computational: the non-nested optima in <<tab:phase_thresholds>> "
                    "and <<tab:greedy>> show that incremental add-one rules used in "
                    "assortment heuristics misrank marginal SKUs for package "
                    "capacities\u2014by "
                    f"{f(metrics['greedy_min_pair_regret'], 1)}\u2013"
                    f"{f(metrics['greedy_max_pair_regret'], 1)} JPY/household-year "
                    f"in {metrics['greedy_fail_count']} of "
                    f"{metrics['greedy_regime_count']} regimes\u2014so the "
                    "inclusion decision requires globally evaluated optima. "
                    "The deterministic-ablation result in Section 3.11 shows this "
                    "non-nestedness is a heterogeneity phenomenon, not a "
                    "randomness phenomenon: the stochastic channels are "
                    "indispensable for mechanism magnitudes (storage, age, "
                    "spoilage windows, adaptation, service) but, in the studied "
                    "calibration, not for the portfolio argmin."
                ),
                {
                    "kind": "subheading",
                    "text": "4.2. Managerial implications",
                },
                paragraph(
                    "Consumption granularity matters when it changes purchase timing or "
                    "inventory exposure enough to activate cost or exit mechanisms. Under "
                    "full carry-over, it is irrelevant to physical discard. When aging "
                    "loss is material, package size changes product loss and freshness; "
                    "when purchase burden dominates, larger packages gain value; when "
                    "packaging dominates, fewer openings and larger sizes gain value."
                ),
                paragraph(
                    "The baseline threshold of "
                    f"{f(metrics['one_to_two_threshold'], 1)} JPY means that a producer "
                    "should add a second size only if annualized manufacturing, "
                    "forecasting, inventory, changeover, and assortment cost per served "
                    "household is below that value. The third-size threshold of "
                    f"{f(metrics['two_to_three_threshold'], 1)} JPY is much tighter. "
                    "These are break-even values for data collection, not estimates of "
                    "actual producer cost. This transparent comparison complements prior "
                    "evidence that variety can increase customer coverage while adding "
                    "inventory, forecasting, manufacturing, and network burdens (Kekre & "
                    "Srinivasan, 1990; Fisher & Ittner, 1999; van Ryzin & Mahajan, 1999; "
                    "Randall & Ulrich, 2001; Wan et al., 2014; Wan & Sanders, 2017; "
                    "Malinowski et al., 2018; Menezes et al., 2021)."
                ),
                paragraph(
                    "Environmental conclusions require product- and material-specific "
                    "life-cycle coefficients; packaging mass and food discard should not "
                    "be collapsed into unsupported carbon values (Fresán et al., 2019; "
                    "Pauer et al., 2019). The modest "
                    "benchmark regret is consistent with current market sizes being "
                    "near-efficient in the modeled scenarios, but revealed market "
                    "optimality cannot be established without prices, shares, promotion, "
                    "channel, and producer-cost data."
                ),
                {
                    "kind": "subheading",
                    "text": "4.3. External validity and manufacturing boundary",
                },
                paragraph(
                    "The theory extends to products with discrete, continuous, or "
                    "count-based use, but the mechanism differs (<<tab:cross_product>>). "
                    "For uncooked rice, carry-over is easy, so mismatch acts through "
                    "adaptation, aging, and storage costs; stable and divisible dry pasta "
                    "normally carries residual inventory over, which suggests a low "
                    "mismatch cost; continuous use of granulated sugar removes the "
                    "lattice mechanism and serves as a conceptual negative control; and "
                    "count-based single-dose medicine can match a prescribed course. "
                    "These examples are theoretical controls and research hypotheses, "
                    "not empirical category findings. A verified 300 g pasta package is "
                    "only a product example (Nisshin Seifun Welna Inc., 2026); its "
                    "event-quantity distribution is not "
                    "estimated here."
                ),
                table("cross_product"),
                paragraph(
                    "Fill-process overage is a separate production lever from nominal "
                    "size: the safety margin in Section 3.10 is set by process precision "
                    "and the underfill constraint, not by consumption granularity, and "
                    "its relative burden falls on smaller packages. Portion-control "
                    "research provides the engineering context (Omar & de Silva, "
                    "2000a, 2000b)."
                ),
                {
                    "kind": "subheading",
                    "text": "4.4. Limitations and empirical agenda",
                },
                paragraph(
                    "The model does not estimate usage-event probabilities, adaptation, "
                    "spoilage, freshness valuation, storage cost, handling cost, purchase "
                    "burden, price–size discounts, unmet-demand penalties, SKU complexity, "
                    "market shares, or population weights. <<tab:parameters>> labels "
                    f"{metrics['unidentified_parameter_count']} inputs as unidentified "
                    f"and {metrics['scenario_parameter_count']} as scenario assumptions. "
                    "Monte Carlo error is not parameter uncertainty."
                ),
                paragraph(
                    "Empirical validation requires linked scanner purchases, household "
                    "inventory diaries, event quantities, lot ages, adaptation, measured "
                    "disposal, prices, and producer assortment costs. A package-size field "
                    "experiment should record carry-over and actual exit rather than infer "
                    "waste from remainders. Producer records should identify whether the "
                    "consumer value of variety exceeds changeover, planning, stocking, and "
                    "distribution complexity."
                ),
            ],
        },
        {
            "heading": "5. Conclusions",
            "blocks": [
                paragraph(
                    "Package–consumption mismatch is not intrinsically wasteful. Under "
                    "full carry-over without an inventory-exit mechanism, physical discard "
                    "is exactly zero regardless of modulo compatibility. Package size "
                    "becomes economically consequential through purchasing, inventory age, "
                    "spoilage, packaging, handling, service, and price mechanisms. Product "
                    "variety is justified only when its globally evaluated consumer benefit "
                    "exceeds SKU complexity under the stated cost model. The rice "
                    "application supplies reproducible scenario evidence and decision "
                    "thresholds; exact portfolio identities remain candidate-grid and "
                    "calibration dependent, and the analysis does not establish a "
                    "Japanese-market optimum. The deterministic ablation bounds "
                    "what the stochastic layer buys: it is required for mechanism "
                    "magnitudes and service/waste pathways, not for the portfolio "
                    "argmin in this calibration. Managers should estimate the "
                    "missing behavioral and producer-cost inputs before converting "
                    "these structural results into package recommendations."
                )
            ],
        },
        {
            "heading": "Declarations",
            "blocks": [
                {"kind": "subheading", "text": "Generative AI declaration"},
                paragraph(
                    "During the preparation of this work the author used Devin "
                    "(devin.ai) in order to format the text and choose words that suited "
                    "the tone, and to help writing codes. After using this tool/service, "
                    "the author reviewed and edited the content as needed and takes full "
                    "responsibility for the content of the published article."
                ),
            ],
        },
    ]


def set_document_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    for name in ["Title", "Heading 1", "Heading 2"]:
        document.styles[name].font.name = "Times New Roman"
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def clear_blinded_metadata(document: Document) -> None:
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""


def add_docx_table(
    document: Document,
    frame: pd.DataFrame,
    key: str,
    number: str | int,
) -> None:
    _, caption_text, note_text = TABLES[key]
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f"Table {number}. {caption_text}")
    run.bold = True
    run.font.size = Pt(10)
    doc_table = document.add_table(rows=1, cols=len(frame.columns))
    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    doc_table.style = "Table Grid"
    for index, column in enumerate(frame.columns):
        cell = doc_table.rows[0].cells[index]
        cell.text = str(column).replace("_", " ")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(7)
            run.font.name = "Times New Roman"
    for values in frame.itertuples(index=False, name=None):
        cells = doc_table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(7)
                run.font.name = "Times New Roman"
    note = document.add_paragraph()
    note.style = document.styles["Normal"]
    note.paragraph_format.space_after = Pt(10)
    if note_text:
        note.add_run(note_text).italic = True


def add_docx_figure(
    document: Document, key: str, number: str | int
) -> None:
    _, filename, caption_text = FIGURES[key]
    path = ROOT / "figures" / filename
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.35))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f"Figure {number}. {caption_text}")
    run.italic = True
    run.font.size = Pt(10)


def add_front_matter(document: Document, abstract: str) -> None:
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    document.add_heading("Abstract", level=1)
    document.add_paragraph(abstract)
    keywords = document.add_paragraph()
    keywords.add_run("Keywords: ").bold = True
    keywords.add_run(
        "package portfolio; SKU proliferation; stochastic inventory simulation; "
        "inventory carry-over; food waste; break-even analysis"
    )


def build_docx(
    sections: list[dict],
    inputs: dict,
    abstract: str,
    path: Path,
    inline_figures: bool,
) -> None:
    document = Document()
    set_document_defaults(document)
    add_front_matter(document, abstract)
    for section in sections:
        document.add_heading(section["heading"], level=1)
        for block in section["blocks"]:
            kind = block["kind"]
            if kind == "subheading":
                document.add_heading(block["text"], level=2)
            elif kind == "paragraph":
                document.add_paragraph(block["text"])
            elif kind == "equation":
                item = document.add_paragraph()
                item.alignment = WD_ALIGN_PARAGRAPH.CENTER
                item.add_run(block["text"]).italic = True
            elif kind == "table":
                add_docx_table(
                    document,
                    table_frame(inputs, block["key"]),
                    block["key"],
                    block["number"],
                )
            elif kind == "figure" and inline_figures:
                add_docx_figure(document, block["key"], block["number"])
    document.add_page_break()
    document.add_heading("References", level=1)
    for reference in REFERENCES:
        item = document.add_paragraph(reference)
        item.paragraph_format.first_line_indent = Inches(-0.25)
        item.paragraph_format.left_indent = Inches(0.25)
        item.paragraph_format.line_spacing = 1.0
    if not inline_figures:
        document.add_page_break()
        document.add_heading("Figure legends", level=1)
        legend_blocks = [
            block
            for section in sections
            for block in section["blocks"]
            if block["kind"] == "figure"
        ]
        for block in legend_blocks:
            caption = FIGURES[block["key"]][2]
            document.add_paragraph(
                f"Figure {block['number']}. {caption}"
            )
    clear_blinded_metadata(document)
    document.save(path)


def frame_to_markdown(frame: pd.DataFrame) -> str:
    clean = frame.copy()
    clean.columns = [column.replace("_", " ") for column in clean.columns]
    header = "| " + " | ".join(clean.columns) + " |"
    separator = "| " + " | ".join("---" for _ in clean.columns) + " |"
    rows = [
        "| "
        + " | ".join(str(value).replace("|", "\\|") for value in values)
        + " |"
        for values in clean.itertuples(index=False, name=None)
    ]
    return "\n".join([header, separator, *rows])


def build_markdown(
    sections: list[dict],
    inputs: dict,
    abstract: str,
) -> str:
    lines = [
        f"# {TITLE}",
        "",
        "## Abstract",
        "",
        abstract,
        "",
        (
            "**Keywords:** package portfolio; SKU proliferation; stochastic "
            "inventory simulation; inventory carry-over; food waste; break-even "
            "analysis"
        ),
        "",
    ]
    for section in sections:
        lines.extend([f"## {section['heading']}", ""])
        for block in section["blocks"]:
            kind = block["kind"]
            if kind == "subheading":
                lines.extend([f"### {block['text']}", ""])
            elif kind == "paragraph":
                lines.extend([block["text"], ""])
            elif kind == "equation":
                lines.extend([f"*{block['text']}*", ""])
            elif kind == "figure":
                number = block["number"]
                _, filename, caption = FIGURES[block["key"]]
                lines.extend(
                    [
                        f"![Figure {number}](../figures/{filename})",
                        "",
                        f"**Figure {number}.** {caption}",
                        "",
                    ]
                )
            elif kind == "table":
                number = block["number"]
                _, caption, note = TABLES[block["key"]]
                lines.extend(
                    [
                        f"**Table {number}. {caption}**",
                        "",
                        frame_to_markdown(table_frame(inputs, block["key"])),
                        "",
                    ]
                )
                if note:
                    lines.extend([f"*{note}*", ""])
    lines.extend(
        [
            "## Supplementary material",
            "",
            (
                "Tables S1-S9 and Figures S1-S4 are provided in the separate "
                "file supplement.docx."
            ),
            "",
            "## References",
            "",
        ]
    )
    for reference in REFERENCES:
        lines.extend([reference, ""])
    return "\n".join(lines)


def build_supplement(
    inputs: dict, supplement_items: list[dict]
) -> None:
    document = Document()
    set_document_defaults(document)
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Supplementary material")
    document.add_paragraph(f"for: {TITLE}")
    document.add_paragraph(
        "This supplement collects the secondary scenario detail, dense-grid "
        "threshold enumerations, robustness rankings, manufacturing-precision "
        "scenarios, and the cross-product illustration referenced in the main "
        "text. All values are generated by the same pipeline as the main "
        "manuscript."
    )
    tables = [item for item in supplement_items if item["kind"] == "table"]
    figures = [item for item in supplement_items if item["kind"] == "figure"]
    markdown_lines = ["# Supplementary material", "", f"for: {TITLE}", ""]
    if tables:
        document.add_heading("Supplementary tables", level=1)
        for item in tables:
            add_docx_table(
                document,
                table_frame(inputs, item["key"]),
                item["key"],
                item["number"],
            )
            _, caption, note = TABLES[item["key"]]
            markdown_lines.extend(
                [
                    f"**Table {item['number']}. {caption}**",
                    "",
                    frame_to_markdown(table_frame(inputs, item["key"])),
                    "",
                ]
            )
            if note:
                markdown_lines.extend([f"*{note}*", ""])
    if figures:
        document.add_heading("Supplementary figures", level=1)
        for item in figures:
            add_docx_figure(document, item["key"], item["number"])
            _, filename, caption = FIGURES[item["key"]]
            markdown_lines.extend(
                [
                    f"![Figure {item['number']}](../figures/{filename})",
                    "",
                    f"**Figure {item['number']}.** {caption}",
                    "",
                ]
            )
    clear_blinded_metadata(document)
    document.save(MANUSCRIPT / "supplement_FINAL.docx")
    (MANUSCRIPT / "supplement_source.md").write_text(
        "\n".join(markdown_lines) + "\n", encoding="utf-8"
    )


def count_words(text: str) -> int:
    return len(re.findall(r"\b[\w–—'-]+\b", text))


def build_title_page(abstract: str, word_count: int) -> None:
    document = Document()
    set_document_defaults(document)
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run("Tatsuki Onishi").bold = True
    affiliation = document.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.add_run(
        "Data Science and AI Innovation Research Promotion Center, Shiga University"
    )
    document.add_heading("Corresponding author", level=1)
    document.add_paragraph(
        "Tatsuki Onishi\nEmail: bougtoir@gmail.com\nORCID: 0000-0001-7261-9062"
    )
    document.add_heading("Article information", level=1)
    document.add_paragraph(
        f"Article type: Research paper\nMain-text word count: {word_count:,}\n"
        f"Abstract word count: {count_words(abstract):,}\n"
        f"Figures: {len(FIGURES) - len(SUPPLEMENT_FIGURES)} "
        f"(+{len(SUPPLEMENT_FIGURES)} supplementary)\n"
        f"Tables: {len(TABLES) - len(SUPPLEMENT_TABLES)} "
        f"(+{len(SUPPLEMENT_TABLES)} supplementary)\n"
        "Supplementary material: supplement.docx "
        f"(Tables S1-S{len(SUPPLEMENT_TABLES)}, "
        f"Figures S1-S{len(SUPPLEMENT_FIGURES)})"
    )
    document.add_heading("CRediT author statement", level=1)
    document.add_paragraph(
        "Tatsuki Onishi: Conceptualization; Methodology; Software; Validation; "
        "Formal analysis; Investigation; Data curation; Writing—original draft; "
        "Writing—review and editing; Visualization; Project administration."
    )
    document.add_heading("Data and code availability", level=1)
    document.add_paragraph(
        "Model code and the generated numerical, figure, and table outputs are "
        f"available at {PUBLIC_REPOSITORY}. Source snapshots restricted to local "
        "research retention are excluded from the public repository; the "
        "publication package reproduces every reported quantity from the "
        "archived inputs and configuration."
    )
    document.add_heading("Declarations", level=1)
    document.add_paragraph(
        "Funding: No external funding.\n"
        "Competing interests: None declared.\n"
        "Ethics: Not applicable; no human participants, personal data, or animals."
    )
    clear_blinded_metadata(document)
    document.save(MANUSCRIPT / "title_page_FINAL.docx")


def build_cover_letter(metrics: dict) -> None:
    document = Document()
    set_document_defaults(document)
    document.add_paragraph("20 October 2026")
    document.add_paragraph(
        "Editor-in-Chief\nSimulation Modelling Practice and Theory"
    )
    document.add_paragraph("Dear Editor,")
    document.add_paragraph(
        f"Please consider my manuscript “{TITLE}” for publication in "
        "Simulation Modelling Practice and Theory."
    )
    document.add_paragraph(
        "Manufacturers repeatedly decide how many package sizes to offer and "
        "which capacities to give them, trading consumer coverage against "
        "changeover, forecasting, stocking, and planning complexity. The paper "
        "submitted here closes a specific gap between two literatures: "
        "package-size optimizers on the supply-chain side (Gámez Albán et al., "
        "2015; Brinker & Gündüz, 2016; Wensing et al., 2018; Zhao et al., 2023) "
        "and product-variety models (Wan et al., 2014; Wan & Sanders, 2017; "
        "Malinowski et al., 2018; Menezes et al., 2021) work under deterministic "
        "or aggregate demand; household waste simulators (Kandemir et al., 2020; "
        "van Rooijen et al., 2024) take the retail offering as fixed; and "
        "Koenigsberg et al. (2010) price a single package size without "
        "downstream inventory states. I couple the two sides by pricing each "
        "candidate consumer SKU inside a stochastic household carry-over "
        "inventory process — FIFO aging, spoilage, backlog, and final-batch "
        "adaptation — and then selecting the offered portfolio by exhaustive "
        "enumeration under an explicit per-SKU complexity charge."
    )
    document.add_paragraph(
        "The paper makes three concrete contributions. First, a zero-discard "
        "boundary lemma: package–consumption mismatch alone cannot create "
        "physical discard without an inventory-exit mechanism — it "
        "establishes a sufficient condition under which arithmetic "
        "residuals cannot generate physical discard. Second, a "
        "global SKU-inclusion rule — an added SKU pays when its consumer value "
        "exceeds its complexity charge — evaluated by exhaustive enumeration "
        "and shown to remain correct when optimal portfolios are not nested, a "
        "case that defeats incremental add-one assortment heuristics\u2014"
        "verified computationally: a nested greedy construction lost "
        f"{f(metrics['greedy_min_pair_regret'], 1)}\u2013"
        f"{f(metrics['greedy_max_pair_regret'], 1)} JPY/household-year at "
        f"the second SKU in {metrics['greedy_fail_count']} of "
        f"{metrics['greedy_regime_count']} prespecified regimes. Third, "
        "matched negative controls and dense phase maps that show consumption "
        "granularity is a second-order term in the studied calibration: under "
        "full carry-over the discrete lattice contributes zero discard and is "
        "second-order in generalized cost. In the "
        f"transparent rice scenarios, the one-to-two and two-to-three SKU "
        "break-even charges were "
        f"{f(metrics['one_to_two_threshold'], 1)} and "
        f"{f(metrics['two_to_three_threshold'], 1)} JPY/household-year per "
        "added SKU — decision thresholds a manager can compare directly with "
        "the firm's own per-SKU cost, not claimed market optima. Package "
        "identities are deliberately reported as calibration-dependent "
        "scenarios rather than Japan-market recommendations."
    )
    document.add_paragraph(
        "I believe this fits SMPT's readership: it applies stochastic "
        "simulation to a management decision problem, makes model development, "
        "implementation, and verification transparent, and contributes three "
        "modelling findings of general interest—a verification boundary the "
        "simulator must reproduce, an experimental-design result that "
        "household heterogeneity alone produces non-nested optimal portfolios "
        "that defeat nested add-one selection rules, and a prespecified "
        "deterministic-counterpart ablation that separates which conclusions "
        "require the stochastic channels (mechanism magnitudes, spoilage "
        "windows, adaptation, backlog) from which do not (the portfolio "
        "argmin in this calibration). Package identities are reported as "
        "calibration-dependent scenarios rather than market optima, and full "
        "code and generated outputs are available in a public GitHub "
        "repository (see the data availability statement)."
    )
    document.add_paragraph(
        "The work is original, is not under consideration elsewhere, and uses "
        "public aggregate sources and simulated scenarios only. A separate "
        "title page, figure files, highlights, supplementary material, and a "
        "reproducibility package accompany the submission."
    )
    document.add_paragraph(
        "Sincerely,\n\nTatsuki Onishi\n"
        "Data Science and AI Innovation Research Promotion Center, Shiga University\n"
        "bougtoir@gmail.com"
    )
    clear_blinded_metadata(document)
    document.save(MANUSCRIPT / "cover_letter_FINAL.docx")


def build_highlights() -> None:
    highlights = [
        "Mismatch alone cannot create discard under full inventory carry-over.",
        "Matched controls separate lattice effects from aging and spoilage.",
        "A global break-even rule determines when an additional SKU pays.",
        "Nested greedy SKU selection loses 106\u2013113 JPY in 6 of 10 regimes.",
        "A deterministic ablation shows which conclusions need stochastic layers.",
    ]
    if any(len(item) > 85 for item in highlights):
        raise ValueError("Each highlight must contain at most 85 characters")
    (MANUSCRIPT / "highlights_FINAL.txt").write_text(
        "\n".join(f"• {item}" for item in highlights) + "\n",
        encoding="utf-8",
    )


def count_main_words(markdown: str) -> int:
    body = markdown.split("## References", maxsplit=1)[0]
    body = re.sub(r"!\[[^\]]+\]\([^)]+\)", "", body)
    body = re.sub(r"\|.*\|", "", body)
    return count_words(body)


def main() -> None:
    MANUSCRIPT.mkdir(parents=True, exist_ok=True)
    inputs = load_inputs()
    metrics = build_revised_metrics(inputs)
    abstract = revised_abstract_text(metrics)
    sections, supplement_items = number_objects(
        build_revised_sections(inputs, metrics)
    )
    markdown = build_markdown(sections, inputs, abstract)
    source_path = MANUSCRIPT / "manuscript_source.md"
    source_path.write_text(markdown, encoding="utf-8")
    build_supplement(inputs, supplement_items)
    build_docx(
        sections,
        inputs,
        abstract,
        MANUSCRIPT / "manuscript_with_figures.docx",
        inline_figures=True,
    )
    build_docx(
        sections,
        inputs,
        abstract,
        MANUSCRIPT / "manuscript_blinded_FINAL.docx",
        inline_figures=False,
    )
    build_title_page(abstract, count_main_words(markdown))
    build_cover_letter(metrics)
    build_highlights()
    figures_dir = MANUSCRIPT / "figures"
    figures_dir.mkdir(parents=True, exist_ok=True)
    for key, (_, filename, _) in FIGURES.items():
        if key in SUPPLEMENT_FIGURES:
            continue
        source = ROOT / "figures" / filename
        if source.exists():
            (figures_dir / filename).write_bytes(source.read_bytes())


if __name__ == "__main__":
    main()
