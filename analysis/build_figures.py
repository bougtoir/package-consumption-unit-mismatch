from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from pptx import Presentation
from pptx.chart.data import ChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "figures"
ORANGE = "#D55E00"
BLUE = "#0072B2"
GREEN = "#009E73"
GRAY = "#6B7280"


def save_figure(figure: plt.Figure, name: str) -> None:
    figure.savefig(FIGURES / f"{name}.png", dpi=300, bbox_inches="tight")
    figure.savefig(FIGURES / f"{name}.pdf", bbox_inches="tight")
    plt.close(figure)


def figure_1_concept() -> None:
    fig, ax = plt.subplots(figsize=(10, 5.2))
    ax.axis("off")
    boxes = [
        (0.03, 0.62, 0.22, 0.22, "Package quantity $P$\npurchases $Q_t$", BLUE),
        (0.39, 0.62, 0.22, 0.22, "Served consumption\n$C_t$", GREEN),
        (0.75, 0.62, 0.22, 0.22, "Carried inventory\n$I_t$", ORANGE),
        (0.15, 0.16, 0.26, 0.22, "Carry-over or\nadaptation", GREEN),
        (0.59, 0.16, 0.26, 0.22, "Inventory-exit\nmechanism $E_t$", ORANGE),
    ]
    for x, y, width, height, text, color in boxes:
        patch = FancyBboxPatch(
            (x, y),
            width,
            height,
            boxstyle="round,pad=0.02",
            edgecolor=color,
            facecolor="white",
            linewidth=2,
        )
        ax.add_patch(patch)
        ax.text(x + width / 2, y + height / 2, text, ha="center", va="center", fontsize=12)
    arrows = [
        ((0.25, 0.73), (0.39, 0.73)),
        ((0.61, 0.73), (0.75, 0.73)),
        ((0.83, 0.62), (0.30, 0.38)),
        ((0.83, 0.62), (0.72, 0.38)),
    ]
    for start, end in arrows:
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="->", mutation_scale=18))
    ax.text(
        0.5,
        0.94,
        "Residual is a state variable; discard requires an explicit mechanism",
        ha="center",
        va="center",
        fontsize=15,
        fontweight="bold",
    )
    save_figure(fig, "figure_1_conceptual_distinction")


def figure_2_lattice() -> None:
    packages = np.linspace(1.5, 10.0, 1200)
    residual = np.mod(packages * 1000.0, 150.0)
    candidates = np.array([1.8, 2.0, 2.1, 2.25, 2.4, 3.0, 4.5, 5.0, 9.0, 10.0])
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.plot(packages, residual, color=BLUE, linewidth=1.5)
    ax.scatter(candidates, np.mod(candidates * 1000.0, 150.0), color=ORANGE, zorder=3)
    ax.set(
        xlabel="Nominal package size (kg)",
        ylabel="Deterministic terminal residual P mod 150 g (g)",
        title="Modulo mismatch is periodic, but it is not a discard measure",
    )
    ax.grid(alpha=0.25)
    save_figure(fig, "figure_2_lattice_residual")


def weighted_size_summary() -> pd.DataFrame:
    frame = pd.read_csv(ROOT / "output/size_household_metrics.csv")
    baseline = frame[frame["regime"] == "baseline"].copy()
    metrics = [
        "generalized_cost_jpy",
        "physical_discard_kg",
        "packaging_mass_kg",
        "purchases",
        "mean_age_at_consumption_days",
    ]
    for metric in metrics:
        baseline[f"weighted_{metric}"] = baseline[metric] * baseline["scenario_weight"]
    return (
        baseline.groupby("package_kg")[
            [f"weighted_{metric}" for metric in metrics]
        ]
        .sum()
        .rename(columns=lambda value: value.replace("weighted_", ""))
        .reset_index()
    )


def figure_3_negative_control() -> None:
    frame = pd.read_csv(ROOT / "output/size_household_metrics.csv")
    selected = frame[
        frame["regime"].isin(
            [
                "baseline",
                "no_discrete_consumption_quantum",
                "randomized_non_lattice_usage",
                "full_carry_no_spoilage",
                "continuous_no_spoilage",
            ]
        )
    ].copy()
    selected["weighted_cost"] = (
        selected["generalized_cost_jpy"] * selected["scenario_weight"]
    )
    selected["weighted_discard"] = (
        selected["physical_discard_kg"] * selected["scenario_weight"]
    )
    summary = selected.groupby(["regime", "package_kg"], as_index=False)[
        ["weighted_cost", "weighted_discard"]
    ].sum()
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
    labels = {
        "baseline": "Discrete + spoilage",
        "no_discrete_consumption_quantum": "Lognormal + spoilage",
        "randomized_non_lattice_usage": "Gamma + spoilage",
        "full_carry_no_spoilage": "Discrete + no spoilage",
        "continuous_no_spoilage": "Lognormal + no spoilage",
    }
    for regime, subset in summary.groupby("regime"):
        axes[0].plot(
            subset["package_kg"],
            subset["weighted_cost"],
            marker="o",
            label=labels[regime],
        )
        axes[1].plot(
            subset["package_kg"],
            subset["weighted_discard"],
            marker="o",
            label=labels[regime],
        )
    axes[0].set(
        xlabel="Package size (kg)",
        ylabel="Generalized cost (JPY/household-year)",
        title="Matched discrete and continuous usage",
    )
    axes[1].set(
        xlabel="Package size (kg)",
        ylabel="Physical discard (kg/household-year)",
        title="No-spoilage controls remain at zero discard",
    )
    for ax in axes:
        ax.grid(alpha=0.25)
    handles, legend_labels = axes[0].get_legend_handles_labels()
    fig.legend(
        handles,
        legend_labels,
        loc="lower center",
        ncol=3,
        bbox_to_anchor=(0.5, -0.06),
    )
    fig.tight_layout()
    save_figure(fig, "figure_3_lattice_negative_control")


def figure_4_consumption_sensitivity() -> None:
    frame = pd.read_csv(ROOT / "output/continuous_consumption_summary.csv")
    frame["best_single_kg"] = frame["best_one_sku_portfolio"].astype(float)
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.8))
    axes[0].step(
        frame["weighted_annual_consumption_kg"],
        frame["best_single_kg"],
        where="mid",
        color=BLUE,
        linewidth=2,
    )
    axes[0].scatter(
        frame["weighted_annual_consumption_kg"],
        frame["best_single_kg"],
        color=BLUE,
    )
    axes[0].set(
        xlabel="Weighted annual consumption (kg/year)",
        ylabel="Optimal single package size (kg)",
        title="Single-size transitions over consumption",
    )
    axes[1].plot(
        frame["weighted_annual_consumption_kg"],
        frame["benchmark_regret_jpy"],
        marker="o",
        color=ORANGE,
    )
    axes[1].set(
        xlabel="Weighted annual consumption (kg/year)",
        ylabel="2|5|10 regret (JPY/household-year)",
        title="Current-size benchmark regret",
    )
    for ax in axes:
        ax.grid(alpha=0.25)
    fig.tight_layout()
    save_figure(fig, "figure_4_consumption_sensitivity")


def pareto_front(frame: pd.DataFrame) -> pd.DataFrame:
    ordered = frame.sort_values(["packaging_mass_kg", "system_cost_jpy"])
    keep = []
    best_cost = np.inf
    for row in ordered.itertuples():
        if row.system_cost_jpy < best_cost:
            keep.append(row.Index)
            best_cost = row.system_cost_jpy
    return frame.loc[keep]


def figure_5_pareto() -> None:
    frame = pd.read_csv(ROOT / "output/portfolio_results.csv")
    baseline = frame[frame["regime"] == "baseline"].copy()
    front = pareto_front(baseline)
    fig, ax = plt.subplots(figsize=(8.5, 5.4))
    for sku_count, subset in baseline.groupby("sku_count"):
        ax.scatter(
            subset["packaging_mass_kg"],
            subset["system_cost_jpy"],
            label=f"{sku_count} SKU",
            alpha=0.55,
        )
    front = front.sort_values("packaging_mass_kg")
    ax.plot(
        front["packaging_mass_kg"],
        front["system_cost_jpy"],
        color="black",
        linewidth=2,
        label="Non-dominated envelope",
    )
    ax.set(
        xlabel="Packaging mass (kg/household-year)",
        ylabel="Generalized system cost (JPY/household-year)",
        title="Portfolio trade-off between packaging and generalized cost",
    )
    ax.legend()
    ax.grid(alpha=0.25)
    save_figure(fig, "figure_5_pareto_frontier")


def figure_6_sku_value() -> None:
    frame = pd.read_csv(ROOT / "output/portfolio_results.csv")
    baseline = frame[frame["regime"] == "baseline"]
    best = baseline.groupby("sku_count", as_index=False)[
        "consumer_cost_jpy"
    ].min()
    complexity = np.linspace(0.0, 600.0, 121)
    fig, ax = plt.subplots(figsize=(8.5, 5))
    for row in best.itertuples(index=False):
        ax.plot(
            complexity,
            row.consumer_cost_jpy + (row.sku_count - 1) * complexity,
            label=f"{row.sku_count} SKU",
        )
    ax.set(
        xlabel="Complexity charge (JPY/household-year per added SKU)",
        ylabel="Minimum system cost (JPY/household-year)",
        title="Global SKU-inclusion thresholds",
    )
    ax.legend()
    ax.grid(alpha=0.25)
    save_figure(fig, "figure_6_sku_break_even")


def figure_7_phase_map() -> None:
    frame = pd.read_csv(ROOT / "output/sku_phase_map.csv")
    hazards = [0.0, 1.0, 4.0]
    fig, axes = plt.subplots(1, 3, figsize=(13, 5.2), sharey=True)
    for ax, hazard in zip(axes, hazards, strict=True):
        subset = frame[frame["spoilage_hazard_multiplier"] == hazard]
        pivot = subset.pivot(
            index="sku_complexity_jpy_per_household_year",
            columns="purchase_event_cost_multiplier",
            values="optimal_sku_count",
        ).sort_index(ascending=True)
        image = ax.imshow(
            pivot.to_numpy(),
            aspect="auto",
            origin="lower",
            cmap="viridis",
            vmin=1,
            vmax=3,
        )
        ax.set_xticks(
            range(len(pivot.columns)),
            [f"{value:g}" for value in pivot.columns],
        )
        ax.set_yticks(
            range(0, len(pivot.index), 4),
            [f"{value:g}" for value in pivot.index[::4]],
        )
        ax.set(
            xlabel="Purchase-event multiplier",
            title=f"Spoilage intensity × {hazard:g}",
        )
    axes[0].set_ylabel("Complexity charge (JPY/household-year/additional SKU)")
    colorbar = fig.colorbar(image, ax=axes, ticks=[1, 2, 3], fraction=0.025)
    colorbar.set_label("Optimal SKU count")
    fig.suptitle("SKU-count phase map across cost and spoilage regimes")
    fig.subplots_adjust(left=0.07, right=0.90, bottom=0.14, top=0.84, wspace=0.18)
    save_figure(fig, "figure_7_sku_phase_map")


def figure_8_candidate_grid() -> None:
    frame = pd.read_csv(ROOT / "output/dense_grid_size_household_metrics.csv")
    frame["weighted_cost"] = (
        frame["generalized_cost_jpy"] * frame["scenario_weight"]
    )
    weighted = frame.groupby("package_kg", as_index=False)["weighted_cost"].sum()
    coarse_sizes = [1.8, 2.0, 2.1, 2.25, 2.4, 3.0, 4.5, 5.0, 9.0, 10.0]
    coarse = weighted[weighted["package_kg"].isin(coarse_sizes)]
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.plot(
        weighted["package_kg"],
        weighted["weighted_cost"],
        color=BLUE,
        linewidth=2,
        label="Dense 0.25 kg grid plus anchors",
    )
    ax.scatter(
        coarse["package_kg"],
        coarse["weighted_cost"],
        color=ORANGE,
        zorder=3,
        label="Original candidate grid",
    )
    dense_best = weighted.loc[weighted["weighted_cost"].idxmin()]
    ax.axvline(dense_best["package_kg"], color=GREEN, linestyle="--")
    ax.set(
        xlabel="Package size (kg)",
        ylabel="Generalized cost (JPY/household-year)",
        title="Candidate-grid robustness of single-size choice",
    )
    ax.legend()
    ax.grid(alpha=0.25)
    save_figure(fig, "figure_8_candidate_grid_robustness")


def figure_9_manufacturing_precision() -> None:
    frame = pd.read_csv(ROOT / "output/manufacturing_precision.csv")
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.8))
    safety = frame.groupby("process_sigma_g", as_index=False)["safety_margin_g"].first()
    axes[0].plot(
        safety["process_sigma_g"],
        safety["safety_margin_g"],
        marker="o",
        color=BLUE,
    )
    axes[0].set(
        xlabel="Fill-process standard deviation (g)",
        ylabel="Safety overfill above nominal target (g)",
        title="Required safety margin",
    )
    for target_g, subset in frame.groupby("target_g"):
        axes[1].plot(
            subset["process_sigma_g"],
            100 * subset["expected_giveaway_fraction"],
            marker="o",
            label=f"{target_g / 1000:g} kg target",
        )
    axes[1].set(
        xlabel="Fill-process standard deviation (g)",
        ylabel="Expected giveaway above nominal (%)",
        title="Smaller packages bear a larger relative giveaway",
    )
    axes[1].legend()
    for ax in axes:
        ax.grid(alpha=0.25)
    fig.suptitle("Manufacturing precision changes the economics of nominal size")
    fig.tight_layout()
    save_figure(fig, "figure_9_manufacturing_precision")


def add_title(slide, title: str) -> None:
    textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(12.2), Inches(0.5))
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(22)
    paragraph.font.bold = True


def add_editable_charts(summary: pd.DataFrame) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 1. Residual inventory is not physical discard")
    box_specs = [
        (0.7, 1.3, "Package quantity P\npurchases Qₜ", BLUE),
        (4.9, 1.3, "Served consumption\nCₜ", GREEN),
        (9.1, 1.3, "Carried inventory\nIₜ", ORANGE),
        (2.5, 4.5, "Carry-over or\nadaptation", GREEN),
        (7.5, 4.5, "Inventory-exit\nmechanism Eₜ", ORANGE),
    ]
    for x, y, text, color in box_specs:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.8), Inches(1.2)
        )
        shape.fill.background()
        shape.line.color.rgb = RGBColor.from_string(color.lstrip("#"))
        shape.text_frame.text = text
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = Pt(18)

    packages = [1.5 + index * 0.05 for index in range(171)]
    data = ChartData()
    data.categories = [f"{value:.2f}" for value in packages]
    data.add_series("P mod 150 g", [value * 1000 % 150 for value in packages])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 2. Deterministic lattice residual")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = False

    controls = pd.read_csv(
        ROOT / "output/mechanism_pathway_comparison.csv"
    )
    controls = controls[
        controls["comparison_basis"] == "baseline_best_single"
    ].copy()
    control_labels = {
        "baseline": "Discrete + spoilage",
        "no_discrete_consumption_quantum": "Lognormal + spoilage",
        "randomized_non_lattice_usage": "Gamma + spoilage",
        "full_carry_no_spoilage": "Discrete + no spoilage",
        "continuous_no_spoilage": "Lognormal + no spoilage",
    }
    data = ChartData()
    data.categories = [
        control_labels[value] for value in controls["regime"]
    ]
    data.add_series("Discard (kg/year)", controls["physical_discard_kg"])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 3. Matched usage and spoilage controls")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    consumption = pd.read_csv(
        ROOT / "output/continuous_consumption_summary.csv"
    )
    data = ChartData()
    data.categories = [
        f"{value:.1f}"
        for value in consumption["weighted_annual_consumption_kg"]
    ]
    data.add_series(
        "System cost (JPY/year)",
        consumption["global_system_cost_jpy"],
    )
    data.add_series(
        "2|5|10 regret (JPY/year)",
        consumption["benchmark_regret_jpy"],
    )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 4. Annual-consumption sensitivity")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    portfolios = pd.read_csv(ROOT / "output/portfolio_results.csv")
    baseline = portfolios[portfolios["regime"] == "baseline"]
    data = XyChartData()
    series = data.add_series("System cost")
    for row in baseline.itertuples(index=False):
        series.add_data_point(row.packaging_mass_kg, row.system_cost_jpy)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 5. Portfolio cost and packaging trade-off")
    slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    )

    thresholds = pd.read_csv(
        ROOT / "output/sku_inclusion_thresholds.csv"
    )
    thresholds = thresholds[thresholds["regime"] == "baseline"].set_index(
        "from_sku_count"
    )
    complexity = list(range(0, 601, 25))
    data = ChartData()
    data.categories = [str(value) for value in complexity]
    for sku_count in [1, 2, 3]:
        if sku_count == 1:
            zero_cost = thresholds.loc[1, "from_consumer_cost_jpy"]
        else:
            zero_cost = thresholds.loc[
                sku_count - 1,
                "to_consumer_cost_jpy",
            ]
        data.add_series(
            f"{sku_count} SKU",
            [zero_cost + (sku_count - 1) * value for value in complexity],
        )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 6. Global SKU break-even thresholds")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    phase = pd.read_csv(ROOT / "output/sku_phase_map.csv")
    phase = phase[
        phase["sku_complexity_jpy_per_household_year"] == 175
    ]
    data = XyChartData()
    for sku_count in [1, 2, 3]:
        series = data.add_series(f"{sku_count} SKU")
        for row in phase[
            phase["optimal_sku_count"] == sku_count
        ].itertuples(index=False):
            series.add_data_point(
                row.purchase_event_cost_multiplier,
                row.spoilage_hazard_multiplier,
            )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 7. SKU-count phase map at 175 JPY complexity")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    grid = pd.read_csv(ROOT / "output/candidate_grid_robustness.csv")
    data = ChartData()
    data.categories = ["1", "2", "3"]
    for grid_name, subset in grid.groupby("candidate_grid"):
        data.add_series(
            grid_name,
            subset.sort_values("sku_count")["system_cost_jpy"],
        )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 8. Candidate-grid robustness")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    manufacturing = pd.read_csv(ROOT / "output/manufacturing_precision.csv")
    data = ChartData()
    data.categories = [
        f"{value:g}"
        for value in sorted(manufacturing["process_sigma_g"].unique())
    ]
    for target_g, subset in manufacturing.groupby("target_g"):
        data.add_series(
            f"{target_g / 1000:g} kg target",
            100
            * subset.sort_values("process_sigma_g")[
                "expected_giveaway_fraction"
            ],
        )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Figure 9. Manufacturing precision")
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.8),
        Inches(1.0),
        Inches(11.8),
        Inches(5.8),
        data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.core_properties.comments = ""
    presentation.save(FIGURES / "pcum_figures_editable.pptx")


def main() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    plt.style.use("seaborn-v0_8-whitegrid")
    summary = weighted_size_summary()
    figure_1_concept()
    figure_2_lattice()
    figure_3_negative_control()
    figure_4_consumption_sensitivity()
    figure_5_pareto()
    figure_6_sku_value()
    figure_7_phase_map()
    figure_8_candidate_grid()
    figure_9_manufacturing_precision()
    add_editable_charts(summary)


if __name__ == "__main__":
    main()
